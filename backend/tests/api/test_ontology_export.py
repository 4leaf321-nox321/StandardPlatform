"""온톨로지 **구조**를 파일로 — 엑셀 여러 표, 또는 다시 넣을 수 있는 JSON.

여기서 지키는 것: 타입마다 표가 하나 서나, 관계와 참조 칸이 함께 나가나, 그리고
**내보낸 JSON 을 그대로 다시 넣을 수 있나**(봉투를 씌우면 가져오기가 거절한다).
"""

from __future__ import annotations

import io
import json
import uuid
from typing import Any

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.modules.ontology import export
from app.modules.workspaces.models import Workspace
from tests.api.conftest import Signed, finish_job
from tests.api.test_ontology import _make_object, _make_property, _make_relation, _make_type


def _world(client: TestClient, admin: Signed) -> tuple[str, str, str]:
    """**이름을 겹치지 않게 짓는다.** 시험 DB 는 스위트가 함께 쓰므로 「부품」 이라는 타입이
    이미 셋 있을 수 있고, 그러면 시트 이름으로 찾은 표가 남의 것이다."""
    tag = uuid.uuid4().hex[:6]
    vendor = _make_type(client, admin, label=f"공급사{tag}", key_policy="optional")
    part = _make_type(client, admin, label=f"부품{tag}", key_policy="optional")
    _make_property(
        client,
        admin,
        part,
        key="grade",
        label="등급",
        data_type="enum",
        enum_options=["A", "B"],
        required=True,
        unit="",
        help="검사 등급",
    )
    _make_property(
        client,
        admin,
        part,
        key="vendor",
        label="공급사",
        data_type="object_ref",
        ref_type_slug=vendor,
    )
    _make_relation(
        client,
        admin,
        "supplied_by",
        label="공급받음",
        src_type_slugs=[part],
        dst_type_slugs=[vendor],
    )
    return vendor, part, f"부품{tag}"


def _book(client: TestClient, who: Signed) -> Any:
    got = client.get("/api/ontology/export?format=xlsx", headers=who.headers)
    assert got.status_code == 200, got.text
    assert "spreadsheetml" in got.headers["content-type"]
    assert ".xlsx" in got.headers["content-disposition"]
    return load_workbook(io.BytesIO(got.content))


def test_타입마다_표가_하나_선다(client: TestClient, admin: Signed) -> None:
    """「각 객체 구조를 개별 테이블로」 — 전체 속성 표만 주면 타입 하나를 보려고
    사람이 엑셀에서 걸러야 한다."""
    _vendor, _part, part_label = _world(client, admin)
    book = _book(client, admin)

    for name in ("개요", "묶음", "타입", "속성", "관계 종류", "참조 칸"):
        assert name in book.sheetnames

    # 시험 DB 는 스위트가 함께 쓴다 — 타입이 쌓여 상한을 넘으면 뒤가 잘리고 이 시험이
    # 무너진다. 그때 고칠 것은 시험이 아니라 상한이다.
    overview = {row[0]: row[1] for row in book["개요"].values}
    assert overview["타입"] <= export.MAX_TYPE_SHEETS, "타입이 상한을 넘었다 — 상한을 올려라"

    assert part_label in book.sheetnames
    rows = [list(row) for row in book[part_label].values]
    assert rows[0][:3] == ["key", "이름", "자료형"]
    keys = {row[0] for row in rows[1:]}
    assert {"grade", "vendor"} <= keys

    grade = next(row for row in rows[1:] if row[0] == "grade")
    assert grade[1] == "등급" and grade[2] == "enum"
    assert grade[3] == "예"  # 필수
    assert grade[6] == "A;B"  # 고를 값


def test_관계와_참조_칸이_함께_나간다(client: TestClient, admin: Signed) -> None:
    """관계 종류만 보면 그림의 절반이다 — 참조 칸도 타입과 타입을 잇는 길이다."""
    vendor, part, _label = _world(client, admin)
    book = _book(client, admin)

    relations = [list(row) for row in book["관계 종류"].values]
    assert any(row[1] == "공급받음" and part in str(row[3]) for row in relations[1:])

    refs = [list(row) for row in book["참조 칸"].values]
    assert any(row[0] == part and row[1] == "vendor" and row[3] == vendor for row in refs[1:])

    # 타입 표에 속성 수·객체 수가 함께 선다 — 「비어 있는 타입」 을 찾는 자리다.
    types = [list(row) for row in book["타입"].values]
    header = types[0]
    line = next(row for row in types[1:] if row[0] == part)
    assert line[header.index("속성 수")] == 2
    assert line[header.index("객체 수")] == 0


def test_내보낸_JSON_은_그대로_다시_넣을_수_있다(client: TestClient, admin: Signed) -> None:
    """봉투를 씌우면 보기에는 친절하지만 가져오기가 모르는 열쇠라며 거절한다."""
    _vendor, part, _label = _world(client, admin)
    got = client.get("/api/ontology/export?format=json", headers=admin.headers)
    assert got.status_code == 200, got.text
    assert ".json" in got.headers["content-disposition"]

    body = json.loads(got.content)
    assert set(body) == {"groups", "types", "relation_types"}
    assert any(one["slug"] == part for one in body["types"])

    again = client.post("/api/ontology/import?dry_run=true", json=body, headers=admin.headers)
    assert again.status_code == 200, again.text
    plan = again.json()
    # 방금 뽑은 것을 그대로 넣으면 **아무것도 안 바뀐다.**
    assert plan["errors"] == []
    assert all(one["action"] == "unchanged" for one in plan["changes"])


def test_읽을_수_있는_사람이면_받는다(client: TestClient, member: Signed) -> None:
    """정의는 비밀이 아니다 — `/ontology/schema` 로 이미 읽는다. 여기만 관리자로 막으면
    같은 것을 보려고 화면을 긁는 길이 생긴다."""
    assert client.get("/api/ontology/export", headers=member.headers).status_code == 200
    assert client.get("/api/ontology/export").status_code == 401


# --- 채워진 것까지 ---------------------------------------------------------------


def _full(client: TestClient, who: Signed, fmt: str) -> Any:
    """작업을 넣고, 워커를 이 프로세스에서 돌리고, 만들어진 파일을 받는다."""
    started = client.post(f"/api/ontology/export?format={fmt}", headers=who.headers)
    assert started.status_code == 202, started.text
    done = finish_job(client, who, started.json())
    assert done["status"] == "done", done
    got = client.get(f"/api/jobs/{done['id']}/download", headers=who.headers)
    assert got.status_code == 200, got.text
    return got


def test_데이터까지_내보내면_타입_시트에_객체_행이_선다(
    client: TestClient, admin: Signed
) -> None:
    """구조만 있는 파일은 「어떤 칸이 있나」 에 답하고, 데이터가 든 파일은 「무엇이 들어
    있나」 에 답한다. 두 물음은 같은 회의에서 함께 나온다."""
    vendor, part, part_label = _world(client, admin)
    acme = _make_object(client, admin, vendor, label="ACME", key="ACME-001")["id"]
    _make_object(
        client, admin, part, label="볼트", key="B-1", properties={"grade": "A", "vendor": acme}
    )

    book = load_workbook(io.BytesIO(_full(client, admin, "xlsx").content))
    assert "관계" in book.sheetnames

    rows = [list(row) for row in book[part_label].values]
    header = rows[0]
    # **열이 일괄 입력이 받는 열 그대로다** — 고쳐서 그대로 다시 붙여 넣을 수 있어야 한다.
    assert header[:3] == ["id", "key", "label"]
    assert "grade" in header and "vendor" in header

    line = next(row for row in rows[1:] if row[header.index("key")] == "B-1")
    assert line[header.index("label")] == "볼트"
    assert line[header.index("grade")] == "A"
    # 참조는 상대의 **식별자**로 — UUID 를 적으면 다른 설치에서 아무것도 안 가리킨다.
    assert line[header.index("vendor")] == "ACME-001"

    overview = {row[0]: row[1] for row in book["개요"].values}
    assert overview["내보낸 객체"] >= 2
    # 시트 이름만 보고는 「속성 정의」 인지 「객체 행」 인지 알 수 없다 — 개요가 말한다.
    assert "객체 행" in str(overview["타입 시트"])
    assert "**" not in str(overview["타입 시트"])  # 엑셀은 마크다운을 안 읽는다


def test_JSON_묶음은_부서마다_나뉘고_그대로_다시_넣을_수_있다(
    client: TestClient, admin: Signed
) -> None:
    """묶음의 한 덩이는 부서 하나를 뜻한다 — 한 덩이로 뭉치면 소유 부서가 사라진다."""
    _vendor, part, _label = _world(client, admin)
    _make_object(client, admin, part, label="너트", key="N-1", properties={"grade": "B"})

    body = json.loads(_full(client, admin, "json").content)
    assert set(body) == {"ontology", "objects", "relations"}

    batch = next(one for one in body["objects"] if one["type_slug"] == part)
    assert batch["workspace_slug"] == admin.workspace
    assert any(row["key"] == "N-1" for row in batch["rows"])
    # id 는 이 설치의 것이다 — 다른 설치에 넣으면 아무것도 안 가리킨다.
    assert all("id" not in row for row in batch["rows"])

    again = client.post("/api/bundles/import", json=body, headers=admin.headers)
    assert again.status_code == 202, again.text
    planned = finish_job(client, admin, again.json())
    assert planned["status"] == "done", planned

    # **내 타입의 행은 그대로 다시 들어간다** — 뽑은 것을 도로 넣으니 바뀔 것이 없다.
    #
    # 묶음 전체의 `ok` 는 보지 않는다: 이 설치에 허브가 관리하는 타입이 있으면 그것은 이쪽에서
    # 못 고치게 막혀 있고(그것이 관리의 뜻이다), 그러면 통째로 다시 넣기는 막힌다. 통째 묶음은
    # **새 설치에 심는 것**이 쓰임새고, 백업으로 읽는 값은 그대로다.
    mine = next(one for one in planned["result"]["objects"] if one["type_slug"] == part)
    assert mine["error"] == ""
    assert mine["plan"]["counts"].get("error", 0) == 0
    assert mine["plan"]["counts"].get("unchanged", 0) >= 1


def test_볼_수_있는_것만_나간다(
    client: TestClient, admin: Signed, member: Signed, db: Session
) -> None:
    """내보내기라고 남의 부서 것이 따라 나가면, 화면에서 막은 것을 파일이 여는 셈이다."""
    _vendor, part, part_label = _world(client, admin)
    other = Workspace(slug=f"other-{uuid.uuid4().hex[:8]}", name="남의팀")
    db.add(other)
    db.commit()
    hidden = client.post(
        f"/api/objects/{part}",
        json={
            "label": "비밀볼트",
            "key": "S-1",
            "workspace_slug": other.slug,
            "properties": {"grade": "A"},
        },
        headers=admin.headers,
    )
    assert hidden.status_code == 201, hidden.text

    book = load_workbook(io.BytesIO(_full(client, member, "xlsx").content))
    keys = {row[1] for row in book[part_label].values}
    assert "S-1" not in keys

    # 시스템 관리자는 전부 본다 — 그래야 통째로 옮길 수 있다.
    everything = load_workbook(io.BytesIO(_full(client, admin, "xlsx").content))
    assert "S-1" in {row[1] for row in everything[part_label].values}
