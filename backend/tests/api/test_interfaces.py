"""인터페이스 — **여러 타입이 한 모양을 따른다**(ADR 0006).

같은 규칙이 두 길(관리 화면 · 정의 가져오기)에서 **같은 말로** 걸리는지를 본다 — 한쪽으로만
막히면, 막힌 사람은 다른 쪽으로 들여보내고 그 상태는 아무 데도 안 뜬다.

시험 DB 는 스위트가 함께 쓰므로 각 시험이 자기 slug 를 만든다.
"""

from __future__ import annotations

import io
from typing import Any

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.modules.ontology.models import OntologySnapshot
from tests.api.conftest import Signed, bundle_export, bundle_import
from tests.api.test_ontology import _import, _make_object, _make_property, _make_type, _uniq

COUNTRY = {
    "key": "country",
    "label": "국가",
    "data_type": "enum",
    "enum_options": ["KR", "US"],
}


def _make_interface(client: TestClient, admin: Signed, base: str = "equip", **kw: Any) -> str:
    slug = kw.pop("slug", None) or _uniq(base)
    body = {"slug": slug, "label": kw.pop("label", base), **kw}
    made = client.post("/api/ontology/interfaces", json=body, headers=admin.headers)
    assert made.status_code == 201, made.text
    return slug


def _add_common(client: TestClient, admin: Signed, iface: str, **kw: Any) -> Any:
    return client.post(
        f"/api/ontology/interfaces/{iface}/properties", json=kw, headers=admin.headers
    )


def _implement(client: TestClient, admin: Signed, type_slug: str, *ifaces: str) -> Any:
    return client.patch(
        f"/api/ontology/types/{type_slug}",
        json={"interface_slugs": list(ifaces)},
        headers=admin.headers,
    )


def _props(client: TestClient, admin: Signed, type_slug: str) -> dict[str, dict[str, Any]]:
    got = client.get(f"/api/ontology/types/{type_slug}/properties", headers=admin.headers)
    assert got.status_code == 200, got.text
    return {one["key"]: one for one in got.json()}


def _code(response: Any) -> str:
    return str(response.json()["error"]["code"])


# --- 구현 ---------------------------------------------------------------------


def test_구현하면_공통_속성이_서고_스키마가_어느_인터페이스의_것인지_적는다(
    client: TestClient, admin: Signed
) -> None:
    iface = _make_interface(client, admin)
    assert _add_common(client, admin, iface, **COUNTRY, required=True).status_code == 201
    kind = _make_type(client, admin, label="시험장비", interface_slugs=[iface])

    props = _props(client, admin, kind)
    assert props["country"]["data_type"] == "enum"
    assert props["country"]["enum_options"] == ["KR", "US"]
    assert props["country"]["required"] is True
    assert props["country"]["interface_slug"] == iface

    schema = client.get("/api/ontology/schema", headers=admin.headers).json()
    found = {one["slug"]: one for one in schema["interfaces"]}[iface]
    assert found["implementers"] == [kind]
    assert [one["key"] for one in found["properties"]] == ["country"]
    types = {one["slug"]: one for one in schema["types"]}
    assert types[kind]["interface_slugs"] == [iface]


def test_같은_모양이면_채택하고_다르면_무엇이_다른지_말하며_거절한다(
    client: TestClient, admin: Signed
) -> None:
    iface = _make_interface(client, admin)
    _add_common(client, admin, iface, **COUNTRY).raise_for_status()

    # 고를 값의 **순서만** 다르면 같은 것이다 — 순서는 인터페이스의 것으로 맞춘다.
    same = _make_type(client, admin, label="계측기")
    _make_property(
        client,
        admin,
        same,
        key="country",
        label="나라",
        data_type="enum",
        enum_options=["US", "KR"],
    )
    assert _implement(client, admin, same, iface).status_code == 200
    adopted = _props(client, admin, same)["country"]
    assert adopted["enum_options"] == ["KR", "US"]
    assert adopted["label"] == "나라"  # 이름은 타입마다다

    # 종류가 다르면 거절 — **화면과 파일이 같은 말을 한다.**
    other = _make_type(client, admin, label="생산설비")
    _make_property(client, admin, other, key="country", label="국가", data_type="text")
    refused = _implement(client, admin, other, iface)
    assert refused.status_code == 409, refused.text
    assert _code(refused).endswith("ONTOLOGY-0007")
    conflicts = refused.json()["error"]["details"]["conflicts"]
    assert any("종류: 인터페이스는 enum, 이 타입은 text" in one for one in conflicts)
    assert _props(client, admin, other)["country"]["interface_slug"] is None

    planned = _import(client, admin, {"types": [{"slug": other, "interface_slugs": [iface]}]})
    assert planned.status_code == 200, planned.text
    assert conflicts[0] in planned.json()["errors"]


def test_구현_미리보기는_아무것도_안_바꾼다(client: TestClient, admin: Signed) -> None:
    iface = _make_interface(client, admin)
    _add_common(client, admin, iface, **COUNTRY).raise_for_status()
    _add_common(
        client, admin, iface, key="maker", label="제조사", data_type="text"
    ).raise_for_status()
    kind = _make_type(client, admin)
    _make_property(client, admin, kind, key="maker", label="제작사", data_type="text")

    plan = client.post(
        f"/api/ontology/types/{kind}/interfaces/plan",
        json={"interface_slugs": [iface]},
        headers=admin.headers,
    )
    assert plan.status_code == 200, plan.text
    body = plan.json()
    assert [one["key"] for one in body["creates"]] == ["country"]
    assert [one["key"] for one in body["adopts"]] == ["maker"]
    assert body["conflicts"] == []
    # 바뀐 것 없음
    assert set(_props(client, admin, kind)) == {"maker"}
    schema = client.get("/api/ontology/schema", headers=admin.headers).json()
    assert {one["slug"]: one for one in schema["types"]}[kind]["interface_slugs"] == []


def test_공통_속성을_고치면_구현_타입이_따라가고_타입에서는_모양을_못_바꾼다(
    client: TestClient, admin: Signed
) -> None:
    iface = _make_interface(client, admin)
    _add_common(client, admin, iface, **COUNTRY).raise_for_status()
    a = _make_type(client, admin, interface_slugs=[iface])
    b = _make_type(client, admin, interface_slugs=[iface])

    changed = client.patch(
        f"/api/ontology/interfaces/{iface}/properties/country",
        json={**COUNTRY, "enum_options": ["KR", "US", "JP"]},
        headers=admin.headers,
    )
    assert changed.status_code == 200, changed.text
    for kind in (a, b):
        assert _props(client, admin, kind)["country"]["enum_options"] == ["KR", "US", "JP"]

    # 타입에서 모양을 바꾸면 거절 — 한 타입만 바뀌면 같은 속성이 타입마다 갈린다.
    body = {**COUNTRY, "enum_options": ["KR"], "label": "국가"}
    locked = client.patch(
        f"/api/ontology/types/{a}/properties/country", json=body, headers=admin.headers
    )
    assert locked.status_code == 409, locked.text
    assert _code(locked).endswith("ONTOLOGY-0008")
    assert locked.json()["error"]["details"]["interface"] == iface
    # 이름 · 묶음은 타입마다 — 된다.
    relabel = client.patch(
        f"/api/ontology/types/{a}/properties/country",
        json={
            **COUNTRY,
            "enum_options": ["KR", "US", "JP"],
            "label": "제조국",
            "section": "출처",
        },
        headers=admin.headers,
    )
    assert relabel.status_code == 200, relabel.text
    assert _props(client, admin, a)["country"]["label"] == "제조국"
    # 지우기 · 고를 값 이름 바꾸기도 인터페이스에서다.
    assert (
        client.delete(
            f"/api/ontology/types/{a}/properties/country", headers=admin.headers
        ).status_code
        == 409
    )
    renamed = client.post(
        f"/api/ontology/types/{a}/properties/country/rename-option",
        json={"from": "KR", "to": "대한민국"},
        headers=admin.headers,
    )
    assert renamed.status_code == 409

    # 구현을 해제하면 속성은 **남고**, 그 뒤로는 타입의 것이다.
    assert _implement(client, admin, a).status_code == 200
    kept = _props(client, admin, a)["country"]
    assert kept["interface_slug"] is None
    free = client.patch(
        f"/api/ontology/types/{a}/properties/country",
        json={**COUNTRY, "enum_options": ["KR"]},
        headers=admin.headers,
    )
    assert free.status_code == 200, free.text


def test_가져오기로_고를_값을_빼면_구현_타입마다_경고한다(
    client: TestClient, admin: Signed
) -> None:
    iface = _make_interface(client, admin)
    _add_common(client, admin, iface, **COUNTRY).raise_for_status()
    kind = _make_type(client, admin, interface_slugs=[iface])
    _make_object(client, admin, kind, label="장비 1", properties={"country": "US"})

    planned = _import(
        client,
        admin,
        {"interfaces": [{"slug": iface, "properties": [{**COUNTRY, "enum_options": ["KR"]}]}]},
    )
    body = planned.json()
    assert body["errors"] == [], body
    assert any(kind in one and "US" in one for one in body["warnings"]), body["warnings"]
    via = [c for c in body["changes"] if c.get("via") == iface]
    assert via and via[0]["slug"] == f"{kind}.country"


def test_구현_타입에_모양이_다른_같은_키가_있으면_공통_속성을_더하지_않는다(
    client: TestClient, admin: Signed
) -> None:
    iface = _make_interface(client, admin)
    kind = _make_type(client, admin, interface_slugs=[iface])
    _make_property(client, admin, kind, key="note", label="비고", data_type="text")
    refused = _add_common(client, admin, iface, key="note", label="비고", data_type="number")
    assert refused.status_code == 409, refused.text
    assert _code(refused).endswith("ONTOLOGY-0007")
    got = client.get(f"/api/ontology/interfaces/{iface}/properties", headers=admin.headers)
    assert got.json() == []


def test_상위_인터페이스를_이어받고_고리와_충돌은_막는다(
    client: TestClient, admin: Signed
) -> None:
    base = _make_interface(client, admin, "asset")
    _add_common(
        client, admin, base, key="maker", label="제조사", data_type="text"
    ).raise_for_status()
    child = _make_interface(client, admin, "equip", extends_slugs=[base])
    kind = _make_type(client, admin, interface_slugs=[child])
    assert _props(client, admin, kind)["maker"]["interface_slug"] == base

    looped = client.patch(
        f"/api/ontology/interfaces/{base}",
        json={"extends_slugs": [child]},
        headers=admin.headers,
    )
    assert looped.status_code == 409, looped.text
    assert _code(looped).endswith("ONTOLOGY-0006")

    # 두 인터페이스가 같은 키를 다른 모양으로 정하면 둘 다 따를 수는 없다.
    other = _make_interface(client, admin, "calib")
    _add_common(
        client, admin, other, key="maker", label="제조사", data_type="number"
    ).raise_for_status()
    clash = _implement(client, admin, kind, child, other)
    assert clash.status_code == 409, clash.text
    assert "다른 모양으로" in clash.json()["error"]["message"]


def test_투영_타입은_구현하지_않고_구현하는_타입은_투영으로_안_바뀐다(
    client: TestClient, admin: Signed
) -> None:
    iface = _make_interface(client, admin)
    made = client.post(
        "/api/ontology/types",
        json={
            "slug": _uniq("dept"),
            "label": "부서",
            "kind_class": "system",
            "system_source": "workspace",
            "interface_slugs": [iface],
        },
        headers=admin.headers,
    )
    assert made.status_code == 409, made.text
    assert _code(made).endswith("ONTOLOGY-0025")

    kind = _make_type(client, admin, interface_slugs=[iface])
    turned = client.patch(
        f"/api/ontology/types/{kind}",
        json={"kind_class": "system", "system_source": "workspace"},
        headers=admin.headers,
    )
    assert turned.status_code == 409, turned.text
    assert _code(turned).endswith("ONTOLOGY-0025")


def test_타입과_인터페이스는_slug_를_함께_쓴다(client: TestClient, admin: Signed) -> None:
    kind = _make_type(client, admin)
    taken = client.post(
        "/api/ontology/interfaces", json={"slug": kind, "label": "x"}, headers=admin.headers
    )
    assert taken.status_code == 409 and _code(taken).endswith("ONTOLOGY-0005")

    iface = _make_interface(client, admin)
    also = client.post(
        "/api/ontology/types", json={"slug": iface, "label": "x"}, headers=admin.headers
    )
    assert also.status_code == 409 and _code(also).endswith("ONTOLOGY-0005")

    same = _uniq("both")
    planned = _import(
        client,
        admin,
        {
            "interfaces": [{"slug": same, "label": "x"}],
            "types": [{"slug": same, "label": "y"}],
        },
    )
    assert planned.json()["errors"], planned.json()


def test_가리키는_것이_있으면_인터페이스를_못_지운다(
    client: TestClient, admin: Signed
) -> None:
    iface = _make_interface(client, admin)
    _add_common(client, admin, iface, **COUNTRY).raise_for_status()
    kind = _make_type(client, admin, interface_slugs=[iface])

    usage = client.get(f"/api/ontology/interfaces/{iface}/usage", headers=admin.headers)
    assert usage.json()["implementers"] == [kind]
    refused = client.delete(f"/api/ontology/interfaces/{iface}", headers=admin.headers)
    assert refused.status_code == 409, refused.text
    assert _code(refused).endswith("ONTOLOGY-0009")
    assert refused.json()["error"]["details"]["implementers"] == [kind]

    _implement(client, admin, kind).raise_for_status()
    assert (
        client.delete(f"/api/ontology/interfaces/{iface}", headers=admin.headers).status_code
        == 204
    )
    assert (
        client.get(
            f"/api/ontology/interfaces/{iface}/properties", headers=admin.headers
        ).status_code
        == 404
    )
    # 구현 타입의 속성은 남는다 — 그 타입의 것이다.
    assert "country" in _props(client, admin, kind)


def test_공통_속성이_될_수_없는_것과_아직_대상이_될_수_없는_것(
    client: TestClient, admin: Signed
) -> None:
    iface = _make_interface(client, admin)
    for body in (
        {"key": "drawing", "label": "도면", "data_type": "file"},
        {"key": "serial", "label": "일련번호", "data_type": "text", "unique": True},
        {"key": "owner", "label": "담당", "data_type": "object_ref"},
        {"key": "grade", "label": "등급", "data_type": "text", "default_value": "A"},
    ):
        refused = _add_common(client, admin, iface, **body)
        assert refused.status_code == 422, (body, refused.text)
        assert _code(refused).endswith("ONTOLOGY-0024")

    # 인터페이스를 참조 대상으로 쓰는 것은 아직 막는다(말하고 막는다). 관계 끝은 된다 —
    # `test_interface_relations.py`.
    kind = _make_type(client, admin)
    ref = client.post(
        f"/api/ontology/types/{kind}/properties",
        json={"key": "eq", "label": "장비", "data_type": "object_ref", "ref_type_slug": iface},
        headers=admin.headers,
    )
    assert ref.status_code == 422 and _code(ref).endswith("ONTOLOGY-0024")


# --- 파일 · 스냅샷 · 묶음 ---------------------------------------------------------


def test_가져오기_한_번으로_인터페이스와_구현이_서고_다시_보내면_그대로다(
    client: TestClient, admin: Signed
) -> None:
    iface, kind = _uniq("equip"), _uniq("tester")
    body = {
        "interfaces": [
            {
                "slug": iface,
                "label": "설비",
                "list_view": {"columns": ["type", "label", "properties.maker"]},
                "properties": [{"key": "maker", "label": "제조사", "data_type": "text"}],
            }
        ],
        "types": [
            {
                "slug": kind,
                "label": "시험장비",
                "interface_slugs": [iface],
                # 구현으로 생기는 속성을 목록이 가리킬 수 있다 — 한 파일 안에서.
                "list_view": {"columns": ["label", "properties.maker"]},
            }
        ],
    }
    applied = _import(client, admin, body, dry_run=False)
    assert applied.status_code == 200 and applied.json()["applied"], applied.text
    assert _props(client, admin, kind)["maker"]["interface_slug"] == iface

    again = _import(client, admin, body)
    touched = [c for c in again.json()["changes"] if c["action"] != "unchanged"]
    assert touched == [], touched

    exported = client.get("/api/ontology/export?format=json", headers=admin.headers).json()
    mine = {one["slug"]: one for one in exported["interfaces"]}[iface]
    assert mine["list_view"]["columns"][0] == "type"
    assert [one["key"] for one in mine["properties"]] == ["maker"]


def test_옛_스냅샷의_상위_타입은_떼고_경고하며_되돌린다(
    client: TestClient, admin: Signed, db: Session
) -> None:
    kind = _make_type(client, admin, label="옛 타입")
    row = OntologySnapshot(
        actor_label="시험",
        reason="인터페이스 전",
        schema={"types": [{"slug": kind, "label": "옛 이름", "parent_slug": "product"}]},
    )
    db.add(row)
    db.commit()

    restored = client.post(f"/api/ontology/snapshots/{row.id}/restore", headers=admin.headers)
    assert restored.status_code == 200, restored.text
    assert restored.json()["applied"] is True
    assert any("상위 타입" in one for one in restored.json()["warnings"])


def test_허브의_인터페이스를_고치면_이_설치의_구현_타입도_따라간다(
    client: TestClient, admin: Signed
) -> None:
    iface = _uniq("hubeq")
    hub = {
        "ontology": {
            "interfaces": [{"slug": iface, "label": "설비", "properties": [COUNTRY]}]
        },
        "source": "hub",
        "apply": True,
    }
    assert bundle_import(client, admin, hub)["applied"] is True

    # 허브의 것은 여기서 못 고친다.
    locked = client.patch(
        f"/api/ontology/interfaces/{iface}", json={"label": "x"}, headers=admin.headers
    )
    assert locked.status_code == 409 and _code(locked).endswith("ONTOLOGY-0082")
    # 이 설치의 타입이 구현하는 것은 된다.
    local = _make_type(client, admin, interface_slugs=[iface])

    grown = {**COUNTRY, "enum_options": ["KR", "US", "JP"]}
    hub["ontology"] = {"interfaces": [{"slug": iface, "properties": [grown]}]}
    assert bundle_import(client, admin, hub)["applied"] is True
    assert _props(client, admin, local)["country"]["enum_options"] == ["KR", "US", "JP"]

    # 이 설치의 구현 타입에 모양이 다른 같은 키가 있으면 묶음 전체가 거절되고, 푸는 법을
    # 말한다.
    _make_property(client, admin, local, key="grade", label="등급", data_type="text")
    hub["ontology"] = {
        "interfaces": [
            {
                "slug": iface,
                "properties": [{"key": "grade", "label": "등급", "data_type": "number"}],
            }
        ]
    }
    refused = bundle_import(client, admin, hub)
    assert refused["applied"] is False
    said = " ".join(refused["ontology"]["errors"])
    assert local in said and "모양" in said


def test_묶음_내보내기는_타입이_구현한_인터페이스를_함께_싣는다(
    client: TestClient, admin: Signed
) -> None:
    group = _uniq("grp")
    client.post(
        "/api/ontology/groups", json={"slug": group, "label": "설비"}, headers=admin.headers
    ).raise_for_status()
    base = _make_interface(client, admin, "asset")
    child = _make_interface(client, admin, "equip", extends_slugs=[base])
    _make_type(client, admin, nav_group_slug=group, interface_slugs=[child])

    bundle = bundle_export(client, admin, group)
    carried = {one["slug"] for one in bundle["ontology"]["interfaces"]}
    assert carried == {base, child}


def test_인터페이스에서_고를_값_이름을_바꾸면_구현_타입_전부의_값이_바뀐다(
    client: TestClient, admin: Signed
) -> None:
    iface = _make_interface(client, admin)
    _add_common(client, admin, iface, **COUNTRY).raise_for_status()
    a = _make_type(client, admin, interface_slugs=[iface])
    b = _make_type(client, admin, interface_slugs=[iface])
    one = _make_object(client, admin, a, label="A1", properties={"country": "KR"})
    two = _make_object(client, admin, b, label="B1", properties={"country": "KR"})

    plan = client.post(
        f"/api/ontology/interfaces/{iface}/properties/country/rename-option",
        json={"from": "KR", "to": "한국"},
        headers=admin.headers,
    )
    assert plan.json()["applied"] is False and plan.json()["objects_with_value"] == 2

    done = client.post(
        f"/api/ontology/interfaces/{iface}/properties/country/rename-option",
        json={"from": "KR", "to": "한국", "apply": True},
        headers=admin.headers,
    )
    assert done.json()["applied"] is True, done.text
    for kind, made in ((a, one), (b, two)):
        got = client.get(f"/api/objects/{kind}/{made['id']}", headers=admin.headers).json()
        assert got["object"]["properties"]["country"] == "한국"
        assert _props(client, admin, kind)["country"]["enum_options"] == ["한국", "US"]


def test_정의_엑셀에_인터페이스와_공통_속성이_실린다(
    client: TestClient, admin: Signed
) -> None:
    """회의에 들고 가는 것은 표다 — 어느 속성이 공통 속성인지 표에서 보여야 한다."""
    iface = _make_interface(client, admin)
    _add_common(client, admin, iface, **COUNTRY).raise_for_status()
    kind = _make_type(client, admin, interface_slugs=[iface])

    got = client.get("/api/ontology/export?format=xlsx", headers=admin.headers)
    assert got.status_code == 200, got.text
    book = load_workbook(io.BytesIO(got.content), read_only=True)
    names = book.sheetnames
    assert names.index("묶음") < names.index("인터페이스") < names.index("타입")

    ifaces = {row[0]: row for row in book["인터페이스"].iter_rows(min_row=2, values_only=True)}
    assert ifaces[iface][3] == kind  # 구현 타입
    common = [row for row in book["공통 속성"].iter_rows(min_row=2, values_only=True)]
    assert (iface, "country") in {(row[0], row[2]) for row in common}

    header = next(book["속성"].iter_rows(max_row=1, values_only=True))
    assert header[-1] == "공통 속성"
    bound = [
        row for row in book["속성"].iter_rows(min_row=2, values_only=True) if row[0] == kind
    ]
    assert [(row[2], row[-1]) for row in bound] == [("country", iface)]
