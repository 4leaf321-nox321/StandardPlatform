"""장표(ADR 0037) — 여러 타입의 기록을 틀 한 장에 모아 본다. 값은 열 때 센다.

    모델 A ← 변경점(모델 · 부품 · 변경 전 · 변경 후)        직접 가리킨다
    모델 A ← 설명회 ─covers→ 잠재리스크(심각도)            설명회를 거쳐서만 이어진다

변경점이 모델을 직접 가리키는지 설명회를 거치는지는 조직마다 다르다(사용자, 2026-10-23) — 둘 다
「주제와 이어지는 칸」 하나로 고른다.
"""

from __future__ import annotations

import uuid
from typing import Any

import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.modules.slides.models import SlideTemplate
from tests.api.conftest import Signed, _login, _make_user
from tests.api.test_ontology import (
    _link,
    _make_object,
    _make_property,
    _make_relation,
    _make_type,
    _uniq,
)


@pytest.fixture
def world(client: TestClient, admin: Signed, member: Signed) -> dict[str, Any]:
    model = _make_type(client, admin, label="모델")
    part = _make_type(client, admin, label="부품")
    change = _make_type(client, admin, label="변경점", usage="log")
    meeting = _make_type(client, admin, label="설명회", usage="log")
    risk = _make_type(client, admin, label="잠재리스크", usage="log")
    for slug in (change, meeting):
        _make_property(
            client,
            admin,
            slug,
            key="model",
            label="모델",
            data_type="object_ref",
            ref_type_slug=model,
        )
    _make_property(
        client,
        admin,
        change,
        key="part",
        label="부품",
        data_type="object_ref",
        ref_type_slug=part,
    )
    _make_property(client, admin, change, key="before", label="변경 전", data_type="text")
    _make_property(client, admin, change, key="after", label="변경 후", data_type="text")
    _make_property(client, admin, risk, key="severity", label="심각도", data_type="text")
    covers = _make_relation(
        client, admin, base="covers", src_type_slugs=[meeting], dst_type_slugs=[risk]
    )

    models = {name: _make_object(client, admin, model, label=name) for name in ("A", "B")}
    connector = _make_object(client, admin, part, label="커넥터")
    changes = {
        name: _make_object(
            client,
            member,
            change,
            label=name,
            properties={
                "model": models[target]["id"],
                "part": connector["id"],
                "before": f"{name} 전",
                "after": f"{name} 후",
            },
        )
        for name, target in (("핀 수", "A"), ("재질", "A"), ("B 변경", "B"))
    }
    talk = _make_object(
        client, member, meeting, label="10월 설명회", properties={"model": models["A"]["id"]}
    )
    risks = {
        name: _make_object(client, member, risk, label=name, properties={"severity": level})
        for name, level in (("접촉 불량", "상"), ("상관없는 것", "하"))
    }
    linked = _link(client, member, meeting, talk["id"], covers, risks["접촉 불량"]["id"])
    assert linked.status_code == 201, linked.text
    # 변경점 「핀 수」 가 일으키는 잠재리스크 둘 — 변경점 → 잠재리스크 관계.
    affects = _make_relation(
        client, admin, base="affects", src_type_slugs=[change], dst_type_slugs=[risk]
    )
    risks["핀 변형"] = _make_object(
        client, member, risk, label="핀 변형", properties={"severity": "중"}
    )
    for name in ("접촉 불량", "핀 변형"):
        made = _link(
            client, member, change, changes["핀 수"]["id"], affects, risks[name]["id"]
        )
        assert made.status_code == 201, made.text
    return {
        "model": model,
        "change": change,
        "risk": risk,
        "covers": covers,
        "affects": affects,
        "parts_type": part,
        "meeting": meeting,
        "talk": talk,
        "models": models,
        "changes": changes,
        "risks": risks,
    }


def _template(world: dict[str, Any], **over: Any) -> dict[str, Any]:
    """좌우(AS-IS / TO-BE — 한 변경점의 두 칸) + 아래 잠재리스크 · 종합 의견."""
    body: dict[str, Any] = {
        "name": _uniq("변경점 검토"),
        "subject_type_slug": world["model"],
        "rows": 2,
        "cols": 2,
        "areas": [
            {
                "key": "asis",
                "title": "AS-IS",
                "row": 0,
                "col": 0,
                "type_slug": world["change"],
                "fields": ["before"],
                "subject_field": "model",
                "ref_line": ["part"],
                "order": "label",
            },
            {
                "key": "tobe",
                "title": "TO-BE",
                "row": 0,
                "col": 1,
                "pair_of": "asis",
                "fields": ["after"],
            },
            {
                "key": "risks",
                "title": "잠재리스크",
                "row": 1,
                "col": 0,
                "type_slug": world["risk"],
                "fields": ["label", "severity"],
                "subject_field": f"in.{world['covers']}.model",
            },
            {"key": "opinion", "title": "종합 의견", "row": 1, "col": 1, "kind": "text"},
        ],
    }
    body.update(over)
    return body


def _made_template(client: TestClient, admin: Signed, world: dict[str, Any]) -> dict[str, Any]:
    made = client.post("/api/slides/templates", json=_template(world), headers=admin.headers)
    assert made.status_code == 201, made.text
    return dict(made.json())


def _slide(
    client: TestClient, who: Signed, template: dict[str, Any], subject: str
) -> dict[str, Any]:
    made = client.post(
        "/api/slides",
        json={"template_id": template["id"], "subject_id": subject},
        headers=who.headers,
    )
    assert made.status_code == 201, made.text
    return dict(made.json())


def _areas(slide: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {one["key"]: one for one in slide["areas"]}


def test_틀은_시스템_관리자만_만든다(
    client: TestClient, member: Signed, world: dict[str, Any]
) -> None:
    got = client.post("/api/slides/templates", json=_template(world), headers=member.headers)
    assert got.status_code == 403


@pytest.mark.parametrize(
    ("change", "says"),
    [
        (lambda b: b["areas"][0].update(fields=["nope"]), "없는 칸"),
        (lambda b: b["areas"][3].update(col=0), "겹칩니다"),
        (lambda b: b["areas"][1].update(row=1, col=1) or b["areas"].pop(3), "맞붙어"),
        (lambda b: b["areas"][0].update(subject_field="ref.nope.model"), "조건"),
        (lambda b: b["areas"][0].update(col_span=3), "격자"),
    ],
)
def test_틀을_저장하기_전에_세워_보고_틀린_것을_말한다(
    client: TestClient,
    admin: Signed,
    world: dict[str, Any],
    change: Any,
    says: str,
) -> None:
    body = _template(world)
    change(body)
    got = client.post("/api/slides/templates", json=body, headers=admin.headers)
    assert got.status_code == 422, got.text
    assert says in got.json()["error"]["message"]


def test_주제의_기록을_모아_그린다_짝_영역은_줄이_맞는다(
    client: TestClient, admin: Signed, member: Signed, world: dict[str, Any]
) -> None:
    slide = _slide(
        client, member, _made_template(client, admin, world), world["models"]["A"]["id"]
    )
    areas = _areas(slide)
    asis, tobe = areas["asis"], areas["tobe"]
    # 모델 A 의 변경점만 — 이름 차례. B 의 변경점은 안 든다.
    assert [one["label"] for one in asis["records"]] == ["재질", "핀 수"]
    assert [one["values"]["before"] for one in asis["records"]] == ["재질 전", "핀 수 전"]
    # 짝 영역은 같은 기록을 같은 차례로, 칸만 다르다.
    assert [one["id"] for one in tobe["records"]] == [one["id"] for one in asis["records"]]
    assert [one["values"]["after"] for one in tobe["records"]] == ["재질 후", "핀 수 후"]
    # 연결 항목 줄 — 기록이 가리키는 부품.
    assert asis["records"][0]["links"] == [
        {
            "label": "부품",
            "items": [
                {
                    "id": asis["records"][0]["links"][0]["items"][0]["id"],
                    "label": "커넥터",
                    "type_slug": asis["records"][0]["links"][0]["items"][0]["type_slug"],
                }
            ],
        }
    ]
    # 설명회를 거쳐서만 이어지는 잠재리스크도 주소 하나로.
    assert [one["label"] for one in areas["risks"]["records"]] == ["접촉 불량"]
    assert areas["risks"]["records"][0]["values"]["severity"] == "상"
    assert areas["opinion"]["kind"] == "text" and areas["opinion"]["text"] == ""
    assert slide["can_edit"] is True and slide["subject"]["label"] == "A"


def test_목록에서_보기는_같은_조건이다(
    client: TestClient, admin: Signed, member: Signed, world: dict[str, Any]
) -> None:
    slide = _slide(
        client, member, _made_template(client, admin, world), world["models"]["A"]["id"]
    )
    for key in ("asis", "risks"):
        area = _areas(slide)[key]
        listed = client.get(
            f"/api/objects/{area['type_slug']}",
            params=area["list_params"],
            headers=member.headers,
        )
        assert listed.status_code == 200, listed.text
        assert listed.json()["total"] == area["total"] == len(area["records"])


def test_살아_있는_값이다_기록이_늘면_장표도_는다(
    client: TestClient, admin: Signed, member: Signed, world: dict[str, Any]
) -> None:
    slide = _slide(
        client, member, _made_template(client, admin, world), world["models"]["A"]["id"]
    )
    _make_object(
        client,
        member,
        world["change"],
        label="도금",
        properties={"model": world["models"]["A"]["id"]},
    )
    again = client.get(f"/api/slides/{slide['id']}", headers=member.headers).json()
    assert [one["label"] for one in _areas(again)["asis"]["records"]] == [
        "도금",
        "재질",
        "핀 수",
    ]


def test_작성자만_고치고_글_영역에만_글을_쓴다(
    client: TestClient,
    db: Session,
    admin: Signed,
    member: Signed,
    world: dict[str, Any],
) -> None:
    slide = _slide(
        client, member, _made_template(client, admin, world), world["models"]["A"]["id"]
    )
    path = f"/api/slides/{slide['id']}"
    written = client.patch(
        path, json={"texts": {"opinion": "접촉 불량 대책 확인 필요"}}, headers=member.headers
    )
    assert written.status_code == 200, written.text
    assert _areas(written.json())["opinion"]["text"] == "접촉 불량 대책 확인 필요"
    wrong = client.patch(path, json={"texts": {"asis": "덮어쓰기"}}, headers=member.headers)
    assert wrong.status_code == 422

    workspace_slug = member.workspace
    from app.modules.workspaces.models import Workspace

    room = db.query(Workspace).filter(Workspace.slug == workspace_slug).one()
    other = _make_user(db, room, label="peer", is_system_admin=False, role="member")
    peer = Signed(email=other.email, token=_login(client, other.email), workspace=room.slug)
    assert (
        client.patch(path, json={"title": "남의 것"}, headers=peer.headers).status_code == 403
    )
    seen = client.get(path, headers=peer.headers)
    assert seen.status_code == 200 and seen.json()["can_edit"] is False


def test_고르기_영역은_고른_차례로_그_타입만(
    client: TestClient, admin: Signed, member: Signed, world: dict[str, Any]
) -> None:
    body = _template(world)
    body["areas"][2] = {
        "key": "risks",
        "title": "잠재리스크",
        "row": 1,
        "col": 0,
        "type_slug": world["risk"],
        "fields": ["severity"],
        "fill": "pick",
    }
    template = client.post("/api/slides/templates", json=body, headers=admin.headers).json()
    slide = _slide(client, member, template, world["models"]["A"]["id"])
    risks = world["risks"]
    picked = [risks["상관없는 것"]["id"], risks["접촉 불량"]["id"]]
    got = client.patch(
        f"/api/slides/{slide['id']}", json={"picks": {"risks": picked}}, headers=member.headers
    )
    assert got.status_code == 200, got.text
    assert [one["id"] for one in _areas(got.json())["risks"]["records"]] == picked
    stranger = client.patch(
        f"/api/slides/{slide['id']}",
        json={"picks": {"risks": [world["changes"]["핀 수"]["id"]]}},
        headers=member.headers,
    )
    assert stranger.status_code == 422


def test_조건이_고장_난_영역만_까닭을_말하고_나머지는_그린다(
    client: TestClient, db: Session, admin: Signed, member: Signed, world: dict[str, Any]
) -> None:
    template = _made_template(client, admin, world)
    slide = _slide(client, member, template, world["models"]["A"]["id"])
    row = db.get(SlideTemplate, uuid.UUID(template["id"]))
    assert row is not None
    areas = [dict(one) for one in row.areas]
    areas[2]["conditions"] = [{"field": "gone", "op": "eq", "value": "x", "exclude": 0}]
    row.areas = areas
    db.commit()
    got = client.get(f"/api/slides/{slide['id']}", headers=member.headers)
    assert got.status_code == 200, got.text
    shown = _areas(got.json())
    assert shown["risks"]["problem"] and not shown["risks"]["records"]
    assert shown["asis"]["problem"] is None and len(shown["asis"]["records"]) == 2


def test_주제는_틀의_주제_타입이어야_한다(
    client: TestClient, admin: Signed, member: Signed, world: dict[str, Any]
) -> None:
    template = _made_template(client, admin, world)
    got = client.post(
        "/api/slides",
        json={"template_id": template["id"], "subject_id": world["changes"]["핀 수"]["id"]},
        headers=member.headers,
    )
    assert got.status_code == 422
    assert "주제" in got.json()["error"]["message"]


def test_장표가_있는_틀과_틀이_쓰는_타입은_못_지운다(
    client: TestClient, admin: Signed, member: Signed, world: dict[str, Any]
) -> None:
    template = _made_template(client, admin, world)
    _slide(client, member, template, world["models"]["A"]["id"])
    held = client.delete(f"/api/slides/templates/{template['id']}", headers=admin.headers)
    assert held.status_code == 409
    # 객체가 하나도 없는 타입이어도 틀의 영역이 쓰면 못 지운다 — FK 가 없어 조용히 깨질 자리다.
    empty = _make_type(client, admin, label="빈 기록", usage="log")
    body = _template(world)
    body["areas"][2].update(type_slug=empty, fields=["label"], subject_field=None)
    used = client.post("/api/slides/templates", json=body, headers=admin.headers)
    assert used.status_code == 201, used.text
    gone = client.delete(f"/api/ontology/types/{empty}", headers=admin.headers)
    assert gone.status_code == 409
    assert "장표의 틀" in gone.text


def test_미리_보기는_저장하지_않은_틀을_세워_본다(
    client: TestClient, admin: Signed, world: dict[str, Any]
) -> None:
    got = client.post(
        "/api/slides/templates/preview",
        json={"template": _template(world), "subject_id": world["models"]["B"]["id"]},
        headers=admin.headers,
    )
    assert got.status_code == 200, got.text
    assert [one["label"] for one in _areas(got.json())["asis"]["records"]] == ["B 변경"]


def _one_change_template(world: dict[str, Any]) -> dict[str, Any]:
    """변경점 하나가 한 장 — AS-IS 에서 하나를 선택하고, 잠재리스크는 그 변경점을 따라온다."""
    body = _template(world)
    body["areas"][0].update(fill="pick", limit=1)
    body["areas"][2] = {
        "key": "risks",
        "title": "잠재리스크",
        "row": 1,
        "col": 0,
        "type_slug": world["risk"],
        "fields": ["label", "severity"],
        "fill": "follow",
        "follow_area": "tobe",  # 짝을 골라도 그 짝(asis)의 기록을 따라간다
        "follow_field": f"in.{world['affects']}",
        "order": "label",
    }
    return body


def test_변경점_하나가_한_장이고_이어진_잠재리스크가_모두_따라온다(
    client: TestClient, admin: Signed, member: Signed, world: dict[str, Any]
) -> None:
    made = client.post(
        "/api/slides/templates", json=_one_change_template(world), headers=admin.headers
    )
    assert made.status_code == 201, made.text
    slide = _slide(client, member, made.json(), world["models"]["A"]["id"])
    path = f"/api/slides/{slide['id']}"
    # 선택 전에는 비어 있고, 따라오는 영역도 비어 있다.
    assert _areas(slide)["asis"]["records"] == [] and _areas(slide)["risks"]["records"] == []

    # 후보는 주제(모델 A)의 변경점만 — B 의 변경점은 후보가 아니다.
    found = client.get(f"{path}/areas/asis/candidates", headers=member.headers)
    assert found.status_code == 200, found.text
    assert {one["label"] for one in found.json()["items"]} == {"핀 수", "재질"}

    changes = world["changes"]
    picked = client.patch(
        path, json={"picks": {"asis": [changes["핀 수"]["id"]]}}, headers=member.headers
    )
    assert picked.status_code == 200, picked.text
    areas = _areas(picked.json())
    assert [one["values"]["before"] for one in areas["asis"]["records"]] == ["핀 수 전"]
    assert [one["values"]["after"] for one in areas["tobe"]["records"]] == ["핀 수 후"]
    # 이어진 잠재리스크는 **모두** — 이름 차례.
    assert [one["label"] for one in areas["risks"]["records"]] == ["접촉 불량", "핀 변형"]
    assert areas["risks"]["follow_title"] == "AS-IS"
    # 「목록에서 보기」 도 같은 수.
    listed = client.get(
        f"/api/objects/{world['risk']}",
        params=areas["risks"]["list_params"],
        headers=member.headers,
    )
    assert listed.status_code == 200 and listed.json()["total"] == 2

    # 하나만 선택하는 영역 · 후보가 아닌 기록은 거절한다.
    two = client.patch(
        path,
        json={"picks": {"asis": [changes["핀 수"]["id"], changes["재질"]["id"]]}},
        headers=member.headers,
    )
    assert two.status_code == 422 and "1건까지" in two.json()["error"]["message"]
    other = client.patch(
        path, json={"picks": {"asis": [changes["B 변경"]["id"]]}}, headers=member.headers
    )
    assert other.status_code == 422 and "후보가 아닌" in other.json()["error"]["message"]


def test_설명회_하나를_선택하면_설명회가_잇는_잠재리스크가_따라온다(
    client: TestClient, admin: Signed, member: Signed, world: dict[str, Any]
) -> None:
    body = _template(world)
    body["areas"] = [
        {
            "key": "talk",
            "title": "설명회",
            "row": 0,
            "col": 0,
            "col_span": 2,
            "type_slug": world["meeting"],
            "fields": ["label"],
            "fill": "pick",
            "limit": 1,
            "subject_field": "model",
        },
        {
            "key": "risks",
            "title": "잠재리스크",
            "row": 1,
            "col": 0,
            "col_span": 2,
            "type_slug": world["risk"],
            "fields": ["label", "severity"],
            "fill": "follow",
            "follow_area": "talk",
            "follow_field": f"in.{world['covers']}",
        },
    ]
    made = client.post("/api/slides/templates", json=body, headers=admin.headers)
    assert made.status_code == 201, made.text
    slide = _slide(client, member, made.json(), world["models"]["A"]["id"])
    got = client.patch(
        f"/api/slides/{slide['id']}",
        json={"picks": {"talk": [world["talk"]["id"]]}},
        headers=member.headers,
    )
    assert got.status_code == 200, got.text
    assert [one["label"] for one in _areas(got.json())["risks"]["records"]] == ["접촉 불량"]


def test_따라가기가_제자리로_돌면_거절한다(
    client: TestClient, admin: Signed, world: dict[str, Any]
) -> None:
    body = _template(world)
    body["areas"][0].update(
        fill="follow", follow_area="risks", follow_field=f"out.{world['affects']}"
    )
    body["areas"][2].update(
        fill="follow", follow_area="asis", follow_field=f"in.{world['affects']}"
    )
    got = client.post("/api/slides/templates", json=body, headers=admin.headers)
    assert got.status_code == 422
    assert "제자리" in got.json()["error"]["message"]


# --- 2단계 — 바로 수정 · 연결 · 새 기록 · 글을 기록으로 · 쓰인 장표 ---------------


def test_기록마다_수정할_수_있나를_객체와_같은_판정으로_말한다(
    client: TestClient, db: Session, admin: Signed, member: Signed, world: dict[str, Any]
) -> None:
    slide = _slide(
        client, member, _made_template(client, admin, world), world["models"]["A"]["id"]
    )
    assert all(one["can_edit"] for one in _areas(slide)["asis"]["records"])

    from app.modules.workspaces.models import Workspace

    room = db.query(Workspace).filter(Workspace.slug == member.workspace).one()
    other = _make_user(db, room, label="peer", is_system_admin=False, role="member")
    peer = Signed(email=other.email, token=_login(client, other.email), workspace=room.slug)
    seen = client.get(f"/api/slides/{slide['id']}", headers=peer.headers).json()
    assert not any(one["can_edit"] for one in _areas(seen)["asis"]["records"])


def test_따라오는_영역을_관계로_잇고_끊으면_장표가_따라온다(
    client: TestClient, admin: Signed, member: Signed, world: dict[str, Any]
) -> None:
    made = client.post(
        "/api/slides/templates", json=_one_change_template(world), headers=admin.headers
    )
    slide = _slide(client, member, made.json(), world["models"]["A"]["id"])
    path = f"/api/slides/{slide['id']}"
    pin = world["changes"]["핀 수"]["id"]
    shown = client.patch(path, json={"picks": {"asis": [pin]}}, headers=member.headers).json()
    risks = _areas(shown)["risks"]
    link = risks["link"]
    assert link["kind"] == "relation" and link["direction"] == "in"
    assert link["relation"] == world["affects"] and link["can_link"] is True
    assert [one["id"] for one in link["sources"]] == [pin]
    # 기록마다 끊을 관계가 실려 온다.
    assert all([tie["source_id"] for tie in one["ties"]] == [pin] for one in risks["records"])

    # 화면이 하는 대로 — 객체 API 로 잇는다(관계는 출발점 = 따라갈 기록에서 맺는다).
    loose = world["risks"]["상관없는 것"]["id"]
    linked = _link(client, member, world["change"], pin, world["affects"], loose)
    assert linked.status_code == 201, linked.text
    again = client.get(path, headers=member.headers).json()
    labels = [one["label"] for one in _areas(again)["risks"]["records"]]
    assert "상관없는 것" in labels

    # 끊으면 빠진다.
    tie = next(
        one["ties"][0]["id"] for one in _areas(again)["risks"]["records"] if one["id"] == loose
    )
    gone = client.delete(
        f"/api/objects/{world['change']}/{pin}/relations/{tie}", headers=member.headers
    )
    assert gone.status_code == 204, gone.text
    after = client.get(path, headers=member.headers).json()
    assert "상관없는 것" not in [one["label"] for one in _areas(after)["risks"]["records"]]


def test_칸으로_이어진_따라오는_영역도_잇는_법을_말한다(
    client: TestClient, admin: Signed, member: Signed, world: dict[str, Any]
) -> None:
    """자기 칸이 가리키는 것(ref)과, 따라갈 기록의 칸이 가리키는 것(backref)."""
    _make_property(
        client,
        admin,
        world["risk"],
        key="change",
        label="원인 변경점",
        data_type="object_ref",
        ref_type_slug=world["change"],
    )
    pin = world["changes"]["핀 수"]["id"]
    patched = client.patch(
        f"/api/objects/{world['risk']}/{world['risks']['접촉 불량']['id']}",
        json={"properties": {"change": pin}},
        headers=member.headers,
    )
    assert patched.status_code == 200, patched.text
    body = _one_change_template(world)
    body["areas"][2]["follow_field"] = "change"
    body["areas"][3] = {
        "key": "parts",
        "title": "부품",
        "row": 1,
        "col": 1,
        "type_slug": world["parts_type"],
        "fields": ["label"],
        "fill": "follow",
        "follow_area": "asis",
        "follow_field": f"in.{world['change']}:part",
    }
    made = client.post("/api/slides/templates", json=body, headers=admin.headers)
    assert made.status_code == 201, made.text
    slide = _slide(client, member, made.json(), world["models"]["A"]["id"])
    shown = client.patch(
        f"/api/slides/{slide['id']}", json={"picks": {"asis": [pin]}}, headers=member.headers
    ).json()
    risks, parts = _areas(shown)["risks"], _areas(shown)["parts"]
    assert risks["link"]["kind"] == "ref" and risks["link"]["property"] == "change"
    assert [one["ties"] for one in risks["records"]] == [[{"id": pin, "source_id": pin}]]
    assert parts["link"]["kind"] == "backref" and parts["link"]["property"] == "part"
    assert [one["label"] for one in parts["records"]] == ["커넥터"]
    assert parts["records"][0]["ties"] == [{"id": pin, "source_id": pin}]


def test_글로_만든_기록을_장표가_기억한다(
    client: TestClient, admin: Signed, member: Signed, world: dict[str, Any]
) -> None:
    slide = _slide(
        client, member, _made_template(client, admin, world), world["models"]["A"]["id"]
    )
    path = f"/api/slides/{slide['id']}"
    risk = _make_object(
        client, member, world["risk"], label="대책 확인 필요", properties={"severity": "중"}
    )
    got = client.patch(path, json={"made": {"opinion": risk["id"]}}, headers=member.headers)
    assert got.status_code == 200, got.text
    assert _areas(got.json())["opinion"]["made"]["label"] == "대책 확인 필요"
    wrong = client.patch(path, json={"made": {"asis": risk["id"]}}, headers=member.headers)
    assert wrong.status_code == 422
    cleared = client.patch(path, json={"made": {"opinion": None}}, headers=member.headers)
    assert _areas(cleared.json())["opinion"]["made"] is None


def test_기록이_쓰인_장표를_말한다(
    client: TestClient, admin: Signed, member: Signed, world: dict[str, Any]
) -> None:
    slide = _slide(
        client, member, _made_template(client, admin, world), world["models"]["A"]["id"]
    )
    pin = world["changes"]["핀 수"]["id"]
    used = client.get(f"/api/slides/used-in/{pin}", headers=member.headers)
    assert used.status_code == 200, used.text
    [one] = used.json()
    assert one["role"] == "record" and one["slide"]["id"] == slide["id"]
    assert sorted(one["areas"]) == ["AS-IS", "TO-BE"]
    subject = client.get(
        f"/api/slides/used-in/{world['models']['A']['id']}", headers=member.headers
    ).json()
    assert [(one["role"], one["slide"]["id"]) for one in subject] == [("subject", slide["id"])]
    # B 의 변경점은 이 장표에 실리지 않았다.
    other = client.get(
        f"/api/slides/used-in/{world['changes']['B 변경']['id']}", headers=member.headers
    ).json()
    assert other == []


# --- 3단계 — pptx 내보내기 -----------------------------------------------------------


def test_pptx_는_영역마다_PPT_표_하나_짝은_한_표(
    client: TestClient, admin: Signed, member: Signed, world: dict[str, Any]
) -> None:
    import io

    from pptx import Presentation
    from pptx.oxml.ns import qn

    slide = _slide(
        client, member, _made_template(client, admin, world), world["models"]["A"]["id"]
    )
    client.patch(
        f"/api/slides/{slide['id']}",
        json={"texts": {"opinion": "접촉 불량 대책 확인\n12월 재검토"}},
        headers=member.headers,
    )
    got = client.get(f"/api/slides/{slide['id']}/pptx", headers=member.headers)
    assert got.status_code == 200, got.text
    assert got.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    assert ".pptx" in got.headers["content-disposition"]

    [page] = Presentation(io.BytesIO(got.content)).slides
    tables = [shape.table for shape in page.shapes if shape.has_table]
    # AS-IS · TO-BE 가 한 표, 잠재리스크 한 표, 종합 의견 한 표.
    assert len(tables) == 3
    pair = tables[0]
    assert [cell.text for cell in pair.rows[0].cells] == ["AS-IS", "TO-BE"]
    line = pair.cell(1, 0)
    assert line.is_merge_origin and line.span_width == 2
    assert line.text == "부품: 커넥터"
    assert [cell.text for cell in pair.rows[2].cells] == ["재질 전", "재질 후"]
    assert [cell.text for cell in pair.rows[4].cells] == ["핀 수 전", "핀 수 후"]
    # 한글 글꼴 · 언어를 박는다 — 안 박으면 PowerPoint 가 영문 글꼴로 그린다.
    props = pair.cell(0, 0).text_frame.paragraphs[0].runs[0]._r.rPr
    assert props.find(qn("a:ea")).get("typeface") == "맑은 고딕"
    assert props.get("lang") == "ko-KR"
    opinion = tables[2]
    assert opinion.cell(0, 0).text == "종합 의견 (작성한 글)"
    assert opinion.cell(1, 0).text == "접촉 불량 대책 확인\n12월 재검토"
    words = " ".join(shape.text_frame.text for shape in page.shapes if shape.has_text_frame)
    assert slide["title"] in words and "기준 시각" in words


def test_pptx_한_장에_안_들면_넘친_기록을_계속_장으로_빠뜨리지_않는다(
    client: TestClient, admin: Signed, member: Signed, world: dict[str, Any]
) -> None:
    import io

    from pptx import Presentation

    for index in range(14):
        _make_object(
            client,
            member,
            world["change"],
            label=f"변경 {index:02d}",
            properties={
                "model": world["models"]["A"]["id"],
                "before": f"변경 {index:02d} 전 — 사양 · 재질 · 치수의 이전 값을 적은 설명",
                "after": f"변경 {index:02d} 후",
            },
        )
    body = _template(world)
    body["areas"][0]["limit"] = 20
    made = client.post("/api/slides/templates", json=body, headers=admin.headers)
    slide = _slide(client, member, made.json(), world["models"]["A"]["id"])
    got = client.get(f"/api/slides/{slide['id']}/pptx", headers=member.headers)
    assert got.status_code == 200, got.text
    pages = list(Presentation(io.BytesIO(got.content)).slides)
    assert len(pages) >= 2
    texts = [
        cell.text
        for page in pages
        for shape in page.shapes
        if shape.has_table
        for row in shape.table.rows
        for cell in row.cells
    ]
    assert any(text.startswith("다음 장에 계속") for text in texts)
    # 기록 열여섯(심은 것 + 핀 수 · 재질)이 모두, 한 번씩 — 빠지거나 겹치지 않는다.
    befores = [text for text in texts if text.endswith(" 전") or "전 — " in text]
    assert len(befores) == 16 and len(set(befores)) == 16
    titles = [
        shape.text_frame.text
        for page in pages[1:]
        for shape in page.shapes
        if shape.has_text_frame
    ]
    assert any("(계속 1/" in one for one in titles)


# --- 4단계 — AI 초안 ----------------------------------------------------------------


def test_AI_가_쓴_글은_초안으로_보이고_사람이_다시_저장하면_풀린다(
    client: TestClient, admin: Signed, member: Signed, world: dict[str, Any]
) -> None:
    slide = _slide(
        client, member, _made_template(client, admin, world), world["models"]["A"]["id"]
    )
    path = f"/api/slides/{slide['id']}"
    drafted = client.patch(
        path,
        json={
            "texts": {"opinion": "핀 수 변경의 접촉 불량(상) 대책 확인 필요"},
            "ai_drafts": ["opinion"],
        },
        headers=member.headers,
    )
    assert drafted.status_code == 200, drafted.text
    assert _areas(drafted.json())["opinion"]["ai_draft"] is True
    human = client.patch(
        path, json={"texts": {"opinion": "확인 완료 — 12월 재검토"}}, headers=member.headers
    )
    assert _areas(human.json())["opinion"]["ai_draft"] is False


def test_같은_장표를_동시에_열어도_부딪히지_않는다(
    client: TestClient, admin: Signed, member: Signed, world: dict[str, Any]
) -> None:
    """열 때마다 「쓰인 장표」 색인을 적는다 — 통째로 지우고 다시 넣던 첫 판은 동시에 연 둘이
    같은 줄을 두 번 넣어 반이 500 이었다(2026-10-24)."""
    from concurrent.futures import ThreadPoolExecutor

    from app.main import app as fastapi_app

    slide = _slide(
        client, member, _made_template(client, admin, world), world["models"]["A"]["id"]
    )

    def one(_: int) -> int:
        with TestClient(fastapi_app, raise_server_exceptions=False) as other:
            got = other.get(f"/api/slides/{slide['id']}", headers=member.headers)
            return int(got.status_code)

    with ThreadPoolExecutor(6) as pool:
        codes = list(pool.map(one, range(12)))
    assert codes == [200] * 12
    used = client.get(
        f"/api/slides/used-in/{world['changes']['핀 수']['id']}", headers=member.headers
    ).json()
    assert [one["slide"]["id"] for one in used] == [slide["id"]]


def test_삭제한_장표만_남은_틀은_지워진다(
    client: TestClient, admin: Signed, member: Signed, world: dict[str, Any]
) -> None:
    """화면은 삭제한 장표를 세지 않아 「0장」 인데 그것들이 틀을 붙들어 삭제가 끝내 안 됐다."""
    template = _made_template(client, admin, world)
    slide = _slide(client, member, template, world["models"]["A"]["id"])
    path = f"/api/slides/templates/{template['id']}"
    assert client.delete(path, headers=admin.headers).status_code == 409
    assert (
        client.delete(f"/api/slides/{slide['id']}", headers=member.headers).status_code == 204
    )
    shown = client.get(path, headers=admin.headers).json()
    assert shown["slides"] == 0
    assert client.delete(path, headers=admin.headers).status_code == 204
    assert client.get(path, headers=admin.headers).status_code == 404


# --- 2026-10-24 점검에서 찾은 것 -------------------------------------------------------


def test_짝이_원본보다_왼쪽이어도_원본으로_그리고_pptx_가_끝난다(
    client: TestClient, admin: Signed, member: Signed, world: dict[str, Any]
) -> None:
    """덩이는 열 차례로 놓인다 — 첫 영역을 원본으로 읽던 때 pptx 쪽 나누기가 끝나지 않았다."""
    import io

    from pptx import Presentation

    for index in range(12):
        _make_object(
            client,
            member,
            world["change"],
            label=f"변경 {index:02d}",
            properties={
                "model": world["models"]["A"]["id"],
                "before": f"변경 {index:02d} 전 — 사양 · 재질 · 치수의 이전 값을 적은 설명",
                "after": f"변경 {index:02d} 후 — 사양 · 재질 · 치수의 새 값을 적은 설명",
            },
        )
    body = _template(world)
    body["areas"][0].update(col=1, limit=20)  # 원본(AS-IS)이 오른쪽
    body["areas"][1].update(col=0)  # 짝(TO-BE)이 왼쪽
    made = client.post("/api/slides/templates", json=body, headers=admin.headers)
    assert made.status_code == 201, made.text
    slide = _slide(client, member, made.json(), world["models"]["A"]["id"])
    got = client.get(f"/api/slides/{slide['id']}/pptx", headers=member.headers)
    assert got.status_code == 200, got.text
    assert len(list(Presentation(io.BytesIO(got.content)).slides)) >= 2


def test_허브가_관리하는_타입_관계는_장표에서_수정_연결을_세우지_않는다(
    client: TestClient, db: Session, admin: Signed, member: Signed, world: dict[str, Any]
) -> None:
    from app.modules.ontology.models import ObjectType, RelationType

    made = client.post(
        "/api/slides/templates", json=_one_change_template(world), headers=admin.headers
    )
    slide = _slide(client, member, made.json(), world["models"]["A"]["id"])
    path = f"/api/slides/{slide['id']}"
    pin = world["changes"]["핀 수"]["id"]
    client.patch(path, json={"picks": {"asis": [pin]}}, headers=member.headers)
    db.query(ObjectType).filter(ObjectType.slug == world["change"]).one().managed_by = "hub"
    db.query(RelationType).filter(
        RelationType.slug == world["affects"]
    ).one().managed_by = "hub"
    db.commit()
    shown = _areas(client.get(path, headers=member.headers).json())
    assert not any(one["can_edit"] for one in shown["asis"]["records"])
    assert shown["risks"]["link"]["can_link"] is False


def test_방향_없는_관계로_따라와도_잇는_방향과_끊을_것을_안다(
    client: TestClient, admin: Signed, member: Signed, world: dict[str, Any]
) -> None:
    """방향이 없는 관계는 양 끝 모두 `out.<관계>` 로 적힌다 — 「out 이면 이 기록이 출발점」
    으로 읽던 때 끊을 것이 비고, 장표에서 맺으면 타입 제한에 걸렸다."""
    related = _make_relation(
        client,
        admin,
        base="related",
        directed=False,
        src_type_slugs=[world["change"]],
        dst_type_slugs=[world["risk"]],
    )
    pin = world["changes"]["핀 수"]["id"]
    linked = _link(
        client, member, world["change"], pin, related, world["risks"]["상관없는 것"]["id"]
    )
    assert linked.status_code == 201, linked.text
    body = _one_change_template(world)
    body["areas"][2]["follow_field"] = f"out.{related}"
    made = client.post("/api/slides/templates", json=body, headers=admin.headers)
    assert made.status_code == 201, made.text
    slide = _slide(client, member, made.json(), world["models"]["A"]["id"])
    shown = client.patch(
        f"/api/slides/{slide['id']}", json={"picks": {"asis": [pin]}}, headers=member.headers
    ).json()
    risks = _areas(shown)["risks"]
    assert [one["label"] for one in risks["records"]] == ["상관없는 것"]
    # 따라갈 기록(변경점)이 출발 타입이라 그쪽에서 맺는다 — 끊을 관계도 실려 온다.
    assert risks["link"]["direction"] == "in"
    assert [len(one["ties"]) for one in risks["records"]] == [1]


def test_주제가_합쳐지면_장표는_이긴_쪽을_주제로_이어_간다(
    client: TestClient, admin: Signed, member: Signed, world: dict[str, Any]
) -> None:
    slide = _slide(
        client, member, _made_template(client, admin, world), world["models"]["A"]["id"]
    )
    merged = client.post(
        f"/api/objects/{world['model']}/{world['models']['A']['id']}/merge",
        json={"into": world["models"]["B"]["id"]},
        headers=admin.headers,
    )
    assert merged.status_code == 200, merged.text
    shown = client.get(f"/api/slides/{slide['id']}", headers=member.headers).json()
    assert shown["subject"]["label"] == "B" and shown["subject"]["deleted"] is False
    # 합친 뒤 변경점은 B 를 가리킨다 — 장표도 그것을 모은다.
    assert {one["label"] for one in _areas(shown)["asis"]["records"]} == {
        "핀 수",
        "재질",
        "B 변경",
    }


def test_칸을_하나도_고르지_않은_영역은_이름을_보인다(
    client: TestClient, admin: Signed, member: Signed, world: dict[str, Any]
) -> None:
    body = _template(world)
    body["areas"][2]["fields"] = []
    made = client.post("/api/slides/templates", json=body, headers=admin.headers)
    assert made.status_code == 201, made.text
    slide = _slide(client, member, made.json(), world["models"]["A"]["id"])
    risks = _areas(slide)["risks"]
    assert risks["fields"] == [{"key": "label", "label": "이름"}]
    assert [one["values"] for one in risks["records"]] == [{"label": "접촉 불량"}]
