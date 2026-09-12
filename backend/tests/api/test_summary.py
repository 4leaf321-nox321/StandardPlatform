"""통계 — **목록과 같은 거르기 위에서 센다.**

따로 세면 「목록에는 12건인데 묶어 보면 15건」 이 되고, 그때 어느 쪽이 맞는지 아무도
모른다. 그래서 이 시험은 언제나 **목록의 total 과 묶음의 합**을 나란히 본다.
"""

from __future__ import annotations

from typing import Any

from fastapi.testclient import TestClient

from tests.api.conftest import Signed
from tests.api.test_ontology import _make_object, _make_property, _make_type

GRADES = ["A", "A", "B", "C", None]


def _parts(client: TestClient, admin: Signed) -> str:
    """등급(고를 값)·무게(숫자)를 가진 타입 하나와 객체 다섯."""
    part = _make_type(client, admin, label="부품")
    _make_property(
        client,
        admin,
        part,
        key="grade",
        label="등급",
        data_type="enum",
        enum_options=["A", "B", "C"],
    )
    _make_property(client, admin, part, key="weight", label="무게", data_type="number")
    for index, grade in enumerate(GRADES, start=1):
        properties: dict[str, Any] = {"weight": index * 10}
        if grade:
            properties["grade"] = grade
        _make_object(client, admin, part, label=f"부품{index}", properties=properties)
    return part


def _summary(client: TestClient, who: Signed, part: str, **params: Any) -> dict[str, Any]:
    got = client.get(f"/api/objects/{part}/summary", params=params, headers=who.headers)
    assert got.status_code == 200, got.text
    return dict(got.json())


def test_속성으로_묶으면_빈_값도_한_칸이_된다(client: TestClient, admin: Signed) -> None:
    part = _parts(client, admin)
    found = _summary(client, admin, part, group_by="properties.grade")

    assert found["group_label"] == "등급"
    assert found["metric_label"] == "건수"
    assert found["total"] == 5
    buckets = {one["label"]: one["count"] for one in found["buckets"]}
    assert buckets == {"A": 2, "B": 1, "C": 1, "(비어 있음)": 1}
    # **막대의 합이 전체와 맞는다.** 빈 값을 숨기면 그 차이가 화면 어디에도 안 적힌다.
    assert sum(one["count"] for one in found["buckets"]) == found["total"]
    assert found["other_groups"] == 0 and found["other_count"] == 0
    # 많은 순으로 — 눈이 맨 위부터 읽는다.
    assert [one["count"] for one in found["buckets"]] == sorted(
        (one["count"] for one in found["buckets"]), reverse=True
    )


def test_거르기를_목록과_같이_받는다(client: TestClient, admin: Signed) -> None:
    part = _parts(client, admin)
    listed = client.get(
        f"/api/objects/{part}", params={"p.grade": "A"}, headers=admin.headers
    ).json()
    found = _summary(client, admin, part, group_by="status", **{"p.grade": "A"})
    assert found["total"] == listed["total"] == 2
    assert [one["label"] for one in found["buckets"]] == ["사용"]


def test_합과_평균은_숫자_칸에서만(client: TestClient, admin: Signed) -> None:
    part = _parts(client, admin)
    found = _summary(
        client,
        admin,
        part,
        group_by="properties.grade",
        metric="sum",
        metric_field="properties.weight",
    )
    values = {one["label"]: one["value"] for one in found["buckets"]}
    assert values["A"] == 30  # 10 + 20
    assert values["(비어 있음)"] == 50
    assert found["metric_label"] == "합계"

    denied = client.get(
        f"/api/objects/{part}/summary",
        params={
            "group_by": "status",
            "metric": "avg",
            "metric_field": "properties.grade",
        },
        headers=admin.headers,
    )
    assert denied.status_code == 422
    assert "숫자 칸이 아니라" in denied.json()["error"]["message"]

    missing = client.get(
        f"/api/objects/{part}/summary",
        params={"group_by": "status", "metric": "sum"},
        headers=admin.headers,
    )
    assert missing.status_code == 422


def test_숫자가_아닌_값이_섞여도_안_죽는다(client: TestClient, admin: Signed) -> None:
    """**행 하나가 질의를 통째로 죽인다** — JSONB 는 글자를 담아 옛 값에 문자가 있을
    수 있고, 그때 화면에는 「합계를 낼 수 없음」 이 아니라 500 이 뜬다."""
    part = _parts(client, admin)
    dirty = _make_object(client, admin, part, label="옛것")
    # 검증을 지나 값이 들어간 상태를 흉내 낸다 — 가져오기 이전의 데이터가 이렇다.
    from app.database import SessionLocal
    from app.modules.objects.models import ObjectInstance

    with SessionLocal() as db:
        row = db.get(ObjectInstance, dirty["id"])
        assert row is not None
        row.properties = {**(row.properties or {}), "weight": "약 3kg"}
        db.commit()

    found = _summary(
        client,
        admin,
        part,
        group_by="status",
        metric="sum",
        metric_field="properties.weight",
    )
    # 못 읽는 값은 셈에서 빠지고, 행 수는 그대로 센다.
    assert found["buckets"][0]["value"] == 150
    assert found["buckets"][0]["count"] == 6


def test_기준으로_쓸_수_없는_칸은_이유를_말한다(client: TestClient, admin: Signed) -> None:
    part = _make_type(client, admin, label="문서")
    _make_property(client, admin, part, key="body", label="본문", data_type="text_long")
    _make_property(
        client, admin, part, key="tags", label="꼬리표", data_type="text", multi=True
    )

    long_text = client.get(
        f"/api/objects/{part}/summary",
        params={"group_by": "properties.body"},
        headers=admin.headers,
    )
    assert long_text.status_code == 422
    assert "그룹이 행 수만큼" in long_text.json()["error"]["message"]

    # **여러 값 칸은 기준이 된다** — 값마다 한 줄로 센다(아래 시험들).
    many = client.get(
        f"/api/objects/{part}/summary",
        params={"group_by": "properties.tags"},
        headers=admin.headers,
    )
    assert many.status_code == 200, many.text

    # 긴 글은 고르개에도 안 뜬다 — 고를 수 없는 것을 보여 주고 나서 거절하지 않는다.
    options = _summary(client, admin, part, group_by="status")["group_options"]
    fields = {one["field"]: one for one in options}
    assert "properties.body" not in fields
    assert fields["properties.tags"]["multi"] is True
    assert {"status", "workspace", "created_year"} <= set(fields)


def test_투영_타입은_여기서_세지_않는다(client: TestClient, admin: Signed) -> None:
    """행이 없다 — 원 표에 물어야 하는 일이고, 그 표의 축을 이 틀은 모른다."""
    projected = _make_type(
        client, admin, label="부서(투영)", kind_class="system", system_source="workspace"
    )
    denied = client.get(f"/api/objects/{projected}/summary", headers=admin.headers)
    assert denied.status_code == 409
    assert "비추는 타입" in denied.json()["error"]["message"]


def _grid(client: TestClient, admin: Signed) -> str:
    """등급(A/B)과 지역(영남/수도권)을 가진 타입 — 기준과 세부 기준으로 나눌 것."""
    part = _make_type(client, admin, label="공급사")
    _make_property(
        client,
        admin,
        part,
        key="grade",
        label="등급",
        data_type="enum",
        enum_options=["A", "B"],
    )
    _make_property(
        client,
        admin,
        part,
        key="area",
        label="지역",
        data_type="enum",
        enum_options=["영남", "수도권"],
    )
    _make_property(client, admin, part, key="score", label="점수", data_type="number")
    rows = [
        ("A", "영남", 10),
        ("A", "영남", 20),
        ("A", "수도권", 30),
        ("B", "수도권", 40),
    ]
    for index, (grade, area, score) in enumerate(rows, start=1):
        _make_object(
            client,
            admin,
            part,
            label=f"공급사{index}",
            properties={"grade": grade, "area": area, "score": score},
        )
    # 지역이 빈 것 하나 — 세부 기준 조각에도 「(비어 있음)」 이 있어야 한다.
    _make_object(client, admin, part, label="공급사5", properties={"grade": "B"})
    return part


def test_세부_기준으로_나누면_조각의_합이_칸과_맞는다(
    client: TestClient, admin: Signed
) -> None:
    """「부서별 몇 건」 다음 물음은 거의 언제나 「그 안에서 등급은」 이다."""
    part = _grid(client, admin)
    found = _summary(
        client, admin, part, group_by="properties.grade", split_by="properties.area"
    )
    assert found["split_label"] == "지역"
    # 계열의 **차례를 서버가 정한다** — 칸마다 나오는 대로 만들면 색이 밀린다.
    assert found["splits"] == ["수도권", "영남", "(비어 있음)"]

    by_label = {one["label"]: one for one in found["buckets"]}
    assert by_label["A"]["count"] == 3
    parts = {one["label"]: one["count"] for one in by_label["A"]["parts"]}
    assert parts == {"영남": 2, "수도권": 1}
    for bucket in found["buckets"]:
        # 조각의 합은 그 칸과 맞는다 — 안 맞으면 그림이 거짓말을 한다.
        assert sum(one["count"] for one in bucket["parts"]) == bucket["count"]
    assert by_label["B"]["parts"][-1]["label"] == "(비어 있음)"


def test_세부_기준이_있어도_합과_평균을_낸다(client: TestClient, admin: Signed) -> None:
    part = _grid(client, admin)
    found = _summary(
        client,
        admin,
        part,
        group_by="properties.grade",
        split_by="properties.area",
        metric="sum",
        metric_field="properties.score",
    )
    a = next(one for one in found["buckets"] if one["label"] == "A")
    values = {one["label"]: one["value"] for one in a["parts"]}
    assert values == {"영남": 30, "수도권": 30}


def test_세부_기준으로_못_쓰는_칸은_저장이_아니라_여기서도_막는다(
    client: TestClient, admin: Signed
) -> None:
    part = _grid(client, admin)
    denied = client.get(
        f"/api/objects/{part}/summary",
        params={"group_by": "properties.grade", "split_by": "properties.note"},
        headers=admin.headers,
    )
    assert denied.status_code == 422


def test_이름으로_묶으면_개별_순위가_된다(client: TestClient, admin: Signed) -> None:
    """**개별 순위가 가장 자주 보는 그림이다.** 「그룹이 행 수만큼」 이라고 막아 두면
    「점수 높은 공급사 열 곳」 을 그릴 방법이 아예 없다 — 상한과 차례가 그것을 순위로
    만든다."""
    part = _grid(client, admin)
    found = _summary(
        client,
        admin,
        part,
        group_by="label",
        metric="sum",
        metric_field="properties.score",
    )
    assert found["group_label"] == "이름"
    assert [one["label"] for one in found["buckets"][:2]] == ["공급사4", "공급사3"]
    assert found["buckets"][0]["value"] == 40

    # 작은 값부터 — 「가장 낮은 것」 을 찾는 물음이 따로 있다.
    lowest = _summary(
        client,
        admin,
        part,
        group_by="label",
        metric="sum",
        metric_field="properties.score",
        order="asc",
    )
    assert lowest["order"] == "asc"
    assert lowest["buckets"][0]["label"] == "공급사1"


def test_식별자를_안_쓰는_타입에는_그_기준이_안_뜬다(
    client: TestClient, admin: Signed
) -> None:
    """고를 수 있다고 보여 주고 나서 빈 그림을 주지 않는다."""
    none_key = _make_type(client, admin, label="메모", key_policy="none")
    fields = {
        one["field"]
        for one in _summary(client, admin, none_key, group_by="label")["group_options"]
    }
    assert "label" in fields and "key" not in fields

    with_key = _make_type(client, admin, label="부품", key_policy="required")
    fields = {
        one["field"]
        for one in _summary(client, admin, with_key, group_by="label")["group_options"]
    }
    assert "key" in fields


def test_모르는_차례는_막는다(client: TestClient, admin: Signed) -> None:
    part = _grid(client, admin)
    denied = client.get(
        f"/api/objects/{part}/summary",
        params={"group_by": "label", "order": "제일 큰 것"},
        headers=admin.headers,
    )
    assert denied.status_code == 422


def _points(client: TestClient, who: Signed, part: str, **params: Any) -> dict[str, Any]:
    got = client.get(f"/api/objects/{part}/points", params=params, headers=who.headers)
    assert got.status_code == 200, got.text
    return dict(got.json())


def test_원값을_그대로_준다_분포는_집계로_안_보인다(client: TestClient, admin: Signed) -> None:
    """평균이 같은 두 공정이 전혀 다른 모양일 수 있고, 그 차이가 대개 문제의 자리다."""
    part = _grid(client, admin)
    found = _points(client, admin, part, x="properties.score", group_by="properties.grade")
    assert found["x_label"] == "점수" and found["group_label"] == "등급"
    assert sorted(one["x"] for one in found["rows"]) == [10, 20, 30, 40]
    groups = {one["group"] for one in found["rows"]}
    assert groups == {"A", "B"}
    assert found["truncated"] is False


def test_값이_없는_행은_점이_안_된다(client: TestClient, admin: Signed) -> None:
    """0 으로 채우면 없는 점이 원점에 모여 그림이 거짓말을 한다."""
    part = _grid(client, admin)
    # 공급사5 는 점수가 없다.
    labels = {
        one["label"] for one in _points(client, admin, part, x="properties.score")["rows"]
    }
    assert "공급사5" not in labels


def test_산점도는_두_칸을_받는다(client: TestClient, admin: Signed) -> None:
    part = _make_type(client, admin, label="시험")
    for key, label in [("load", "하중"), ("stress", "응력")]:
        _make_property(client, admin, part, key=key, label=label, data_type="number")
    _make_object(client, admin, part, label="시험1", properties={"load": 10, "stress": 100})
    found = _points(client, admin, part, x="properties.load", y="properties.stress")
    assert found["y_label"] == "응력"
    assert found["rows"][0]["x"] == 10 and found["rows"][0]["y"] == 100


def test_숫자가_아닌_칸으로는_못_그린다(client: TestClient, admin: Signed) -> None:
    """고를 수 있다고 보여 주고 나서 빈 그림을 주지 않는다."""
    part = _grid(client, admin)
    denied = client.get(
        f"/api/objects/{part}/points",
        params={"x": "properties.grade"},
        headers=admin.headers,
    )
    assert denied.status_code == 422
    assert "숫자 칸이 아니라" in denied.json()["error"]["message"]


# --- 파일로 내보내기 -----------------------------------------------------------------


def test_통계_표를_엑셀로_내보낸다_화면과_같은_숫자로(
    client: TestClient, admin: Signed
) -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    part = _parts(client, admin)
    got = client.get(
        f"/api/objects/{part}/summary/export",
        params={"group_by": "properties.grade", "format": "xlsx"},
        headers=admin.headers,
    )
    assert got.status_code == 200, got.text
    assert "spreadsheetml" in got.headers["content-type"]
    sheet = load_workbook(BytesIO(got.content)).active
    table = [[cell.value for cell in row] for row in sheet.iter_rows()]
    assert table[0] == ["등급", "건수"]
    assert {row[0]: row[1] for row in table[1:]} == {
        "A": 2,
        "B": 1,
        "C": 1,
        "(비어 있음)": 1,
        "전체": 5,
    }


def test_세부_기준이_있으면_그_값이_열이_되고_CSV_는_BOM_이_붙는다(
    client: TestClient, admin: Signed
) -> None:
    import csv
    import io

    part = _parts(client, admin)
    got = client.get(
        f"/api/objects/{part}/summary/export",
        params={"group_by": "status", "split_by": "properties.grade", "format": "csv"},
        headers=admin.headers,
    )
    assert got.status_code == 200, got.text
    # 엑셀이 한글을 안 깨뜨리게.
    assert got.content.startswith("﻿".encode())
    rows = list(csv.reader(io.StringIO(got.content.decode("utf-8-sig"))))
    assert rows[0][0] == "상태" and rows[0][-1] == "합계"
    assert set(rows[0][1:-1]) == {"A", "B", "C", "(비어 있음)"}
    assert rows[-1][0] == "전체" and rows[-1][-1] == "5"


def test_수식으로_읽히는_이름은_글자로_막는다(client: TestClient, admin: Signed) -> None:
    """이름은 사람이 적는다 — `=HYPERLINK(...)` 가 셀에서 실행되면 안 된다."""
    import csv
    import io

    part = _parts(client, admin)
    _make_object(client, admin, part, label='=HYPERLINK("http://x")', properties={"weight": 1})
    got = client.get(
        f"/api/objects/{part}/summary/export",
        params={"group_by": "label", "format": "csv"},
        headers=admin.headers,
    )
    labels = [row[0] for row in csv.reader(io.StringIO(got.content.decode("utf-8-sig")))]
    assert '\'=HYPERLINK("http://x")' in labels
    assert '=HYPERLINK("http://x")' not in labels


def test_원값도_파일로_나간다(client: TestClient, admin: Signed) -> None:
    import csv
    import io

    part = _parts(client, admin)
    got = client.get(
        f"/api/objects/{part}/points/export",
        params={"x": "properties.weight", "group_by": "properties.grade", "format": "csv"},
        headers=admin.headers,
    )
    assert got.status_code == 200, got.text
    rows = list(csv.reader(io.StringIO(got.content.decode("utf-8-sig"))))
    assert rows[0] == ["이름", "무게", "등급"]
    assert len(rows) == 1 + 5


# --- 여러 값 칸을 기준으로 ---------------------------------------------------------

#: 해석 분야(여러 값)와 출시 연도를 가진 툴 넷. 하나는 분야가 비어 있다.
TOOLS = [
    ("툴1", ["구조", "유체"], 2000),
    ("툴2", ["구조"], 2010),
    ("툴3", [], 2020),
    ("툴4", ["유체", "열"], 2016),
]


def _tools(client: TestClient, admin: Signed) -> str:
    tool = _make_type(client, admin, label="툴")
    _make_property(
        client,
        admin,
        tool,
        key="field",
        label="해석 분야",
        data_type="enum",
        enum_options=["구조", "유체", "열"],
        multi=True,
    )
    _make_property(client, admin, tool, key="year", label="출시", data_type="number")
    for label, fields, year in TOOLS:
        properties: dict[str, Any] = {"year": year}
        if fields:
            properties["field"] = fields
        _make_object(client, admin, tool, label=label, properties=properties)
    return tool


def test_여러_값_칸은_값마다_세고_합이_전체보다_클_수_있다고_알린다(
    client: TestClient, admin: Signed
) -> None:
    """「해석 분야별 툴 수」 — 막아 두면 사람이 실제로 묻는 그림이 아예 안 나온다."""
    tool = _tools(client, admin)
    found = _summary(client, admin, tool, group_by="properties.field")

    buckets = {one["label"]: one["count"] for one in found["buckets"]}
    assert buckets == {"구조": 2, "유체": 2, "열": 1, "(비어 있음)": 1}
    # **전체는 객체 수다.** 막대의 합(6)이 전체(4)보다 크고, 그 사실을 알린다.
    assert found["total"] == 4
    assert found["overlap"] is True
    assert found["other_count"] == 0 and found["other_groups"] == 0

    # 막대의 수와 그 값으로 거른 목록의 수가 같다 — 막대를 누르면 그대로 걸러진다.
    listed = client.get(
        f"/api/objects/{tool}", params={"f.field.eq": "구조"}, headers=admin.headers
    ).json()
    assert listed["total"] == buckets["구조"]

    # 여러 값이 아닌 기준이면 알리지 않는다.
    assert _summary(client, admin, tool, group_by="status")["overlap"] is False


def test_여러_값_칸으로_평균을_내고_세부_기준으로도_쓴다(
    client: TestClient, admin: Signed
) -> None:
    tool = _tools(client, admin)
    averaged = _summary(
        client,
        admin,
        tool,
        group_by="properties.field",
        metric="avg",
        metric_field="properties.year",
    )
    values = {one["label"]: one["value"] for one in averaged["buckets"]}
    assert values["구조"] == 2005 and values["열"] == 2016

    split = _summary(client, admin, tool, group_by="status", split_by="properties.field")
    assert split["overlap"] is True
    active = split["buckets"][0]
    parts = {one["label"]: one["count"] for one in active["parts"]}
    # **빈 값도 조각이다** — 분야가 없는 툴3 이 「(비어 있음)」 조각으로 선다.
    assert parts == {"구조": 2, "유체": 2, "열": 1, "(비어 있음)": 1}


def test_여러_값_기준으로_나누면_비어_있음_칸에도_조각이_선다(
    client: TestClient, admin: Signed
) -> None:
    """IN 에는 NULL 이 안 걸린다 — 따로 안 적으면 「(비어 있음)」 칸만 조각이 빈다."""
    tool = _tools(client, admin)
    found = _summary(client, admin, tool, group_by="properties.field", split_by="status")
    empty = next(one for one in found["buckets"] if one["key"] is None)
    assert [part["count"] for part in empty["parts"]] == [1]


def test_여러_값_기준의_파일은_전체가_객체_수라고_적는다(
    client: TestClient, admin: Signed
) -> None:
    import csv
    import io

    tool = _tools(client, admin)
    got = client.get(
        f"/api/objects/{tool}/summary/export",
        params={"group_by": "properties.field", "format": "csv"},
        headers=admin.headers,
    )
    rows = list(csv.reader(io.StringIO(got.content.decode("utf-8-sig"))))
    assert rows[-1] == ["전체 (객체 수)", "4"]


def test_여러_값_기준으로_상자를_가르면_값마다_점이_든다(
    client: TestClient, admin: Signed
) -> None:
    tool = _tools(client, admin)
    got = client.get(
        f"/api/objects/{tool}/points",
        params={"x": "properties.year", "group_by": "properties.field"},
        headers=admin.headers,
    )
    assert got.status_code == 200, got.text
    groups = sorted(one["group"] for one in got.json()["rows"])
    assert groups == ["(비어 있음)", "구조", "구조", "열", "유체", "유체"]
