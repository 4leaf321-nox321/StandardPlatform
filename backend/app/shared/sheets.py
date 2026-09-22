"""표를 파일로 — CSV 와 Excel.

어느 모듈이든 「이 표를 내려받게」 가 필요해지면 여기를 쓴다. 화면마다 따로 만들면
BOM·수식 막기·파일 이름 규칙이 조금씩 어긋나고, 그 어긋남은 **엑셀에서 한글이 깨지는
날** 드러난다.

## 수식으로 읽히는 글자를 막는다

객체 이름은 사람이 적는 값이다. `=HYPERLINK(...)` 나 `+cmd|...` 로 시작하는 이름이
그대로 셀에 들어가면 **엑셀이 그것을 수식으로 실행한다**(CSV injection). 글자로 시작이
`= + - @` 이거나 탭·줄바꿈이면 앞에 `'` 를 붙여 글자로 둔다. 숫자는 건드리지 않는다.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import UTC, datetime
from typing import Any

from fastapi import Response

#: 이 글자로 시작하는 **글자** 값은 엑셀이 수식으로 읽는다.
FORMULA_LEADS = ("=", "+", "-", "@", "\t", "\r")

CSV_TYPE = "text/csv; charset=utf-8"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def safe_cell(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(FORMULA_LEADS):
        return "'" + value
    return value


def to_csv(header: list[str], rows: list[list[Any]]) -> bytes:
    """엑셀이 한글을 안 깨뜨리게 BOM 을 붙인다."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([safe_cell(one) for one in header])
    for row in rows:
        writer.writerow(["" if one is None else safe_cell(one) for one in row])
    return ("\ufeff" + buffer.getvalue()).encode("utf-8")


@dataclass
class Page:
    """워크북의 시트 하나 — 이름과 표."""

    name: str
    header: list[str] = field(default_factory=list)
    rows: list[list[Any]] = field(default_factory=list)


def sheet_name(raw: str, taken: set[str]) -> str:
    """엑셀이 받는 시트 이름 — 31자까지, 몇몇 글자는 못 쓰고, **겹치면 안 된다.**

    겹치는 이름을 그대로 주면 openpyxl 이 조용히 바꾸거나 터진다. 여기서 번호를 붙여
    **무엇이 잘렸는지 보이게** 한다.
    """
    cleaned = "".join(ch for ch in raw if ch not in "[]:*?/\\").strip("'") or "표"
    name = cleaned[:31]
    serial = 2
    while name.lower() in taken:
        suffix = f" ({serial})"
        name = cleaned[: 31 - len(suffix)] + suffix
        serial += 1
    taken.add(name.lower())
    return name


def to_workbook(pages: list[Page]) -> bytes:
    """여러 표를 한 파일로 — **시트마다 하나.** 머리줄은 굵게, 첫 줄은 고정.

    한 파일에 담는 이유: 표 여섯 개를 파일 여섯 개로 주면 받는 사람이 그것을 다시
    한 곳에 모아야 하고, 그 사이에 하나가 빠진다.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    book.remove(book.active)
    taken: set[str] = set()
    for one in pages or [Page("표")]:
        page = book.create_sheet(sheet_name(one.name, taken))
        page.append([safe_cell(cell) for cell in one.header])
        for cell in page[1]:
            cell.font = Font(bold=True)
        for row in one.rows:
            page.append(["" if cell is None else safe_cell(cell) for cell in row])
        page.freeze_panes = "A2"
        for index, title in enumerate(one.header, start=1):
            widest = max(
                [
                    len(str(title)),
                    *(len(str(r[index - 1])) for r in one.rows if len(r) >= index),
                ]
            )
            page.column_dimensions[page.cell(row=1, column=index).column_letter].width = min(
                60, max(8, widest * 1.6)
            )
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def to_xlsx(header: list[str], rows: list[list[Any]], *, sheet: str = "표") -> bytes:
    return to_workbook([Page(sheet, header, rows)])


def stamped(stem: str, ext: str) -> str:
    """파일 이름은 **ASCII** 로 — 헤더에 한글을 그대로 넣으면 깨지는 프록시가 있다.
    사람이 읽는 이름은 화면이 붙인다."""
    return f"{stem}-{datetime.now(UTC).strftime('%Y%m%d-%H%M')}.{ext}"


def attachment(body: bytes, *, media: str, stem: str, ext: str) -> Response:
    return Response(
        body,
        media_type=media,
        headers={"Content-Disposition": f'attachment; filename="{stamped(stem, ext)}"'},
    )


def workbook_response(pages: list[Page], *, stem: str) -> Response:
    """여러 표를 엑셀 파일 하나로 내려준다."""
    return attachment(to_workbook(pages), media=XLSX_TYPE, stem=stem, ext="xlsx")


def file_response(
    header: list[str], rows: list[list[Any]], *, fmt: str, stem: str, sheet: str = "표"
) -> Response:
    """내려받기 응답. 파일 이름은 **ASCII** 로 — 헤더에 한글을 그대로 넣으면 깨지는
    프록시가 있다. 사람이 읽는 이름은 화면이 붙인다."""
    if fmt == "xlsx":
        return attachment(
            to_xlsx(header, rows, sheet=sheet), media=XLSX_TYPE, stem=stem, ext="xlsx"
        )
    return attachment(to_csv(header, rows), media=CSV_TYPE, stem=stem, ext="csv")
