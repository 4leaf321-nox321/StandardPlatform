"""장표를 pptx 로(ADR 0037 3단계) — 화면과 같은 모양: 격자의 영역마다 **PPT 표 하나**.

장표는 살아 있는 값이고, 그날의 모습은 이 파일이 남긴다 — 그래서 기준 시각을 제목 아래와 바닥에
찍는다. 표는 화면의 표와 같은 열을 쓴다(`SlideGrid.columnsOf`): 영역마다 너비 몫을 「칸 이름 |
값」 두 열(칸이 하나면 값 한 열)로 나누고, 머리글 · 연결 항목 줄 · 칸 줄이 같은 열을 쓴다. 짝
영역(AS-IS / TO-BE)은 표 하나다.

**높이는 내용이 정한다.** 격자의 행을 똑같이 나누면 기록이 많은 영역이 아래 영역을 덮는다(첫
판에서 그랬다). 그래서 표마다 줄 수 · 글 길이로 높이를 어림해 격자 행의 높이를 정하고, 한 장을
넘으면 글자를 함께 줄이고(8pt 까지), 그래도 넘치면 넘친 기록을 「(계속)」 장으로 넘긴다 —
빠뜨리지 않는다. python-pptx 는 **여기서만** 부른다(requirements 의 주석).
"""

from __future__ import annotations

import math
from dataclasses import dataclass
from io import BytesIO
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

from app.modules.slides.schemas import AreaOut, RecordOut, SlideOut

WIDTH_IN = 13.333
HEIGHT_IN = 7.5
MARGIN_IN = 0.4
TOP_IN = 1.15
BOTTOM_IN = 0.45
GAP_IN = 0.12
FONT = "맑은 고딕"
LABEL_SHARE = 0.32
#: 글자 크기(pt) — 머리글 · 값 · 칸 이름과 연결 항목 줄. 한 장을 넘으면 함께 줄인다.
HEAD_PT, VALUE_PT, SMALL_PT = 11.0, 10.0, 9.0
MIN_SCALE = 0.8


def _rgb(hex_value: str) -> RGBColor:
    """python-pptx 의 색은 타입이 없다 — 여기 한 곳에서만 부른다."""
    found: RGBColor = RGBColor.from_string(hex_value)  # type: ignore[no-untyped-call]
    return found


HEAD = _rgb("E5E7EB")
LINE = _rgb("F3F4F6")
WHITE = _rgb("FFFFFF")
MUTED = _rgb("6B7280")
#: 「No Style, Table Grid」 — 가는 검은 테두리, 칠 없음. 기본 표 모양(파란 칠)은 장표에 안
#: 맞는다.
GRID_STYLE = "{5940675A-B579-460E-94D1-54222C63F5DA}"


@dataclass
class _Block:
    group: list[AreaOut]
    left: float
    width: float
    heights: list[float]
    """표의 줄마다 높이(인치) — 만들 때 이 높이로 줄을 세운다."""


@dataclass
class _Page:
    """pptx 한 장 — 격자 · 영역(기록이 잘렸을 수 있다) · 다음 장으로 넘긴 기록 수."""

    rows: int
    cols: int
    areas: list[AreaOut]
    continued: dict[str, int]
    """덩이(첫 영역의 키)마다 다음 장으로 넘긴 기록 수."""
    suffix: str = ""


def build(slide: SlideOut) -> bytes:
    """장표 한 장 → pptx. 한 장에 안 들면(8pt 까지 줄여도) 넘친 기록을 「(계속)」 장으로 —
    빠뜨리지 않는다."""
    deck = Presentation()
    deck.slide_width = Inches(WIDTH_IN)
    deck.slide_height = Inches(HEIGHT_IN)
    shown = slide.rendered_at.astimezone().strftime("%Y-%m-%d %H:%M")
    subject = slide.subject.label if slide.subject else "(없음)"
    inner = WIDTH_IN - 2 * MARGIN_IN
    for spec in _paginate(slide):
        page = deck.slides.add_slide(deck.slide_layouts[6])  # 빈 모양
        _box(page, MARGIN_IN, 0.3, inner, 0.5, slide.title + spec.suffix, 22, bold=True)
        _box(
            page,
            MARGIN_IN,
            0.75,
            inner,
            0.3,
            f"주제: {subject} · 틀: {slide.template_name} · 기준 시각: {shown}",
            11,
            color=MUTED,
        )
        scale, blocks, rows = _fit(spec)
        for block in blocks:
            first = block.group[0]
            top = TOP_IN + sum(rows[: first.row]) + GAP_IN * first.row
            if first.kind == "text":
                _text_table(page, first, block, top, scale)
            else:
                continued = spec.continued.get(primary_of(block.group).key, 0)
                _records_table(page, block, top, scale, continued)
        _box(
            page,
            MARGIN_IN,
            HEIGHT_IN - 0.4,
            inner,
            0.3,
            f"기준 시각 {shown} — 장표는 열 때마다 지금의 기록을 읽습니다. "
            "이 파일은 그때의 사본입니다.",
            9,
            color=MUTED,
        )
    out = BytesIO()
    deck.save(out)
    return out.getvalue()


def _room(rows: int) -> float:
    return HEIGHT_IN - TOP_IN - BOTTOM_IN - GAP_IN * (rows - 1)


def _fit(page: _Page) -> tuple[float, list[_Block], list[float]]:
    """한 장에 들 때까지 글자를 줄여 본다(`MIN_SCALE` 까지) — 줄이면 한 줄에 드는 글자가 늘어
    줄 수도 준다."""
    scale = 1.0
    blocks, rows = _layout(page, scale)
    for _ in range(5):
        if sum(rows) <= _room(page.rows) or scale <= MIN_SCALE:
            break
        scale = max(MIN_SCALE, scale * _room(page.rows) / sum(rows) * 0.97)
        blocks, rows = _layout(page, scale)
    return scale, blocks, rows


def _fits(page: _Page) -> bool:
    return sum(_fit(page)[2]) <= _room(page.rows)


def _paginate(slide: SlideOut) -> list[_Page]:
    """첫 장은 틀 그대로. 안 들면 가장 긴 덩이의 끝 기록부터 다음 장으로 넘긴다 — 넘친 기록은
    그 덩이만 담은 「(계속)」 장에, 한 장에 들 만큼씩."""
    moved: dict[str, int] = {}

    def trimmed() -> _Page:
        areas: list[AreaOut] = []
        for one in slide.areas:
            cut = moved.get(one.pair_of or one.key, 0)
            if cut and one.kind == "records":
                one = one.model_copy(update={"records": one.records[: len(one.records) - cut]})
            areas.append(one)
        return _Page(slide.rows, slide.cols, areas, dict(moved))

    first = trimmed()
    while not _fits(first):
        longest = [
            group
            for group in blocks_of(first.areas)
            if group[0].kind == "records" and len(group[0].records) > 1
        ]
        if not longest:
            break
        group = max(
            longest,
            key=lambda one: len(one[0].records) * max(1, *(len(area.fields) for area in one)),
        )
        key = primary_of(group).key
        moved[key] = moved.get(key, 0) + 1
        first = trimmed()
    pages = [first]
    place = {one.key: (one.row, one.col) for one in slide.areas}
    # 넘긴 덩이는 격자의 차례로(위에서 아래, 왼쪽에서 오른쪽) — 넘긴 차례가 아니라.
    for key, count in sorted(moved.items(), key=lambda one: place[one[0]]):
        group = sorted(
            (one for one in slide.areas if one.key == key or one.pair_of == key),
            key=lambda one: one.col,
        )
        pending = list(range(len(group[0].records) - count, len(group[0].records)))
        while pending:
            take = len(pending)
            while take > 1 and not _fits(
                _rest(group, pending[:take], last=take == len(pending))
            ):
                take -= 1
            pages.append(_rest(group, pending[:take], last=take == len(pending)))
            pending = pending[take:]
    for number, page in enumerate(pages[1:], 1):
        page.suffix = f" (계속 {number}/{len(pages) - 1})"
    return pages


def _rest(group: list[AreaOut], indices: list[int], *, last: bool) -> _Page:
    """넘친 기록만 담은 장 — 덩이를 한 행 전체 너비로. 마지막 장에만 「외 N건」(상한 너머)이
    남는다."""
    areas: list[AreaOut] = []
    col = 0
    for one in group:
        records = [one.records[index] for index in indices if index < len(one.records)]
        beyond = one.total - len(one.records) if last else 0
        areas.append(
            one.model_copy(
                update={
                    "row": 0,
                    "row_span": 1,
                    "col": col,
                    "title": f"{one.title} (계속)",
                    "records": records,
                    "total": len(records) + max(0, beyond),
                }
            )
        )
        col += one.col_span
    return _Page(1, col, areas, {})


def blocks_of(areas: list[AreaOut]) -> list[list[AreaOut]]:
    """짝끼리 묶은 덩이 — 화면의 `blocksOf` 와 같다. 왼쪽부터."""
    out: list[list[AreaOut]] = []
    for area in areas:
        if area.pair_of:
            continue
        group = [area, *(one for one in areas if one.pair_of == area.key)]
        out.append(sorted(group, key=lambda one: one.col))
    return out


def primary_of(group: list[AreaOut]) -> AreaOut:
    """덩이의 원본 — 짝이 아닌 영역. 덩이는 열 차례로 놓여 짝이 왼쪽에 올 수 있다 — 그때 첫
    영역을 원본으로 읽으면 넘긴 기록을 짝의 키로 세어 쪽 나누기가 끝나지 않았다(2026-10-24
    점검)."""
    return next((one for one in group if not one.pair_of), group[0])


def columns_of(group: list[AreaOut]) -> list[tuple[AreaOut, str, float]]:
    """덩이의 열 — (영역, `label`|`value`, 너비 몫). 화면의 `columnsOf` 와 같다."""
    span = sum(one.col_span for one in group)
    out: list[tuple[AreaOut, str, float]] = []
    for one in group:
        share = one.col_span / span
        if len(one.fields) > 1:
            out.append((one, "label", share * LABEL_SHARE))
            out.append((one, "value", share * (1 - LABEL_SHARE)))
        else:
            out.append((one, "value", share))
    return out


# --- 높이 어림 ------------------------------------------------------------------


def _layout(slide: _Page, scale: float) -> tuple[list[_Block], list[float]]:
    """덩이마다 너비 · 줄 높이, 격자 행마다 높이. 여러 행에 걸친 덩이는 모자란 만큼 마지막 행에
    더한다."""
    cell = (WIDTH_IN - 2 * MARGIN_IN - GAP_IN * (slide.cols - 1)) / slide.cols
    blocks: list[_Block] = []
    for group in blocks_of(slide.areas):
        first, last = group[0], group[-1]
        span = last.col + last.col_span - first.col
        width = span * cell + (span - 1) * GAP_IN
        heights = (
            _text_heights(first, width, scale)
            if first.kind == "text"
            else _records_heights(
                group, width, scale, slide.continued.get(primary_of(group).key, 0)
            )
        )
        blocks.append(_Block(group, MARGIN_IN + first.col * (cell + GAP_IN), width, heights))
    rows = [0.0] * slide.rows
    for block in blocks:
        first = block.group[0]
        if first.row_span == 1:
            rows[first.row] = max(rows[first.row], sum(block.heights))
    for block in blocks:
        first = block.group[0]
        if first.row_span > 1:
            spanned = range(first.row, first.row + first.row_span)
            have = sum(rows[one] for one in spanned) + GAP_IN * (first.row_span - 1)
            rows[first.row + first.row_span - 1] += max(0.0, sum(block.heights) - have)
    return blocks, rows


def _units(text: str) -> float:
    """글의 폭(em) — 한글 · 한자는 1, 그 밖은 0.55 로 어림한다."""
    return sum(1.0 if ord(ch) >= 0x2E80 else 0.55 for ch in text)


def _lines(text: str, width: float, size: float) -> int:
    per_line = max(width - 0.12, 0.3) * 72 / size
    return sum(max(1, math.ceil(_units(line) / per_line)) for line in text.split("\n"))


def _height(lines: int, size: float) -> float:
    return lines * size * 1.25 / 72 + 0.08


def _records_heights(
    group: list[AreaOut], width: float, scale: float, continued: int
) -> list[float]:
    """`_records_table` 이 세우는 줄과 같은 차례로 — 머리글, 기록마다 연결 항목 줄과 칸 줄,
    끝의 알림."""
    head, value_pt, small = HEAD_PT * scale, VALUE_PT * scale, SMALL_PT * scale
    primary = primary_of(group)
    shares = [share * width for _, _, share in columns_of(group)]
    depth_of = max(1, *(len(one.fields) for one in group))
    out = [max(_height(_lines(one.title, width * 0.5, head), head) for one in group)]
    for index, record in enumerate(primary.records):
        out.append(_height(_lines(_line(record), width, small), small))
        rows = [_height(1, value_pt)] * depth_of
        merged: list[float] = []
        at = 0
        for one in group:
            shown = one.records[index] if index < len(one.records) else None
            if len(one.fields) > 1:
                for depth, field in enumerate(one.fields):
                    value = (shown.values.get(field.key) or "—") if shown else ""
                    need = max(
                        _height(_lines(field.label, shares[at], small), small),
                        _height(_lines(value, shares[at + 1], value_pt), value_pt),
                    )
                    rows[depth] = max(rows[depth], need)
                at += 2
                continue
            only = one.fields[0] if one.fields else None
            value = (shown.values.get(only.key) or "—") if shown and only else ""
            merged.append(_height(_lines(value, shares[at], value_pt), value_pt))
            at += 1
        for need in merged:
            if sum(rows) < need:
                rows[-1] += need - sum(rows)
        out.extend(rows)
    note = _note(primary, continued)
    if note:
        out.append(_height(_lines(note, width, small), small))
    return out


def _text_heights(area: AreaOut, width: float, scale: float) -> list[float]:
    head, value_pt = HEAD_PT * scale, VALUE_PT * scale
    return [
        _height(1, head),
        _height(_lines(area.text or "—", width, value_pt), value_pt),
    ]


# --- 표 ----------------------------------------------------------------------


def _records_table(page: Any, block: _Block, top: float, scale: float, continued: int) -> None:
    group = block.group
    primary = primary_of(group)
    records = primary.records
    columns = columns_of(group)
    depth_of = max(1, *(len(one.fields) for one in group))
    note = _note(primary, continued)
    table = _table(page, block, top)
    for index, (_, _, share) in enumerate(columns):
        table.columns[index].width = Inches(block.width * share)
    head, value_pt, small = HEAD_PT * scale, VALUE_PT * scale, SMALL_PT * scale

    # 머리글 — 영역마다 그 열들을 합친다.
    at = 0
    for one in group:
        cols = 2 if len(one.fields) > 1 else 1
        cell = table.cell(0, at)
        if cols == 2:
            cell.merge(table.cell(0, at + 1))
        _write(cell, one.title, head, bold=True, fill=HEAD)
        at += cols

    row = 1
    for index, record in enumerate(records):
        line = table.cell(row, 0)
        if len(columns) > 1:
            line.merge(table.cell(row, len(columns) - 1))
        _write(line, _line(record), small, fill=LINE, color=MUTED)
        row += 1
        for depth in range(depth_of):
            at = 0
            for one in group:
                shown = one.records[index] if index < len(one.records) else None
                field = one.fields[depth] if depth < len(one.fields) else None
                value = (shown.values.get(field.key) or "—") if shown and field else ""
                if len(one.fields) > 1:
                    label = field.label if field else ""
                    _write(table.cell(row + depth, at), label, small, color=MUTED)
                    _write(table.cell(row + depth, at + 1), value, value_pt)
                    at += 2
                    continue
                if depth == 0:
                    cell = table.cell(row + depth, at)
                    if depth_of > 1:
                        cell.merge(table.cell(row + depth_of - 1, at))
                    _write(cell, value, value_pt)
                at += 1
        row += depth_of
    if note:
        cell = table.cell(row, 0)
        if len(columns) > 1:
            cell.merge(table.cell(row, len(columns) - 1))
        _write(cell, note, small, color=MUTED)


def _text_table(page: Any, area: AreaOut, block: _Block, top: float, scale: float) -> None:
    """글 영역 — 머리글과 글 한 칸. 사람이 쓴 글이라 머리에 그렇다고 적는다."""
    table = _table(page, block, top)
    table.columns[0].width = Inches(block.width)
    _write(
        table.cell(0, 0),
        f"{area.title} ({'작성한 글 · AI 초안' if area.ai_draft else '작성한 글'})",
        HEAD_PT * scale,
        bold=True,
        fill=HEAD,
    )
    _write(table.cell(1, 0), area.text or "—", VALUE_PT * scale)


def _note(area: AreaOut, continued: int = 0) -> str | None:
    if continued:
        return f"다음 장에 계속 — {continued}건"
    if area.problem:
        return f"이 영역을 표시하지 못했습니다 — {area.problem}"
    if not area.records:
        if area.fill == "pick":
            return "선택한 기록이 없습니다."
        if area.fill == "follow":
            return f"「{area.follow_title or ''}」 의 기록과 이어진 기록이 없습니다."
        return "해당하는 기록이 없습니다."
    more = area.total - len(area.records)
    return (
        f"외 {more}건 — 장표 화면의 「목록에서 보기」 에서 확인합니다." if more > 0 else None
    )


def _line(record: RecordOut) -> str:
    """연결 항목 줄 — 「부품: 커넥터 · 모델: A」. 없으면 기록 이름."""
    parts = [
        f"{one.label}: {', '.join(item.label for item in one.items)}"
        for one in record.links
        if one.items
    ]
    return " · ".join(parts) if parts else record.label


def _table(page: Any, block: _Block, top: float) -> Any:
    rows = len(block.heights)
    cols = len(columns_of(block.group)) if block.group[0].kind == "records" else 1
    shape = page.shapes.add_table(
        rows,
        cols,
        Inches(block.left),
        Inches(top),
        Inches(block.width),
        Inches(sum(block.heights)),
    )
    table = shape.table
    table.first_row = False
    table.horz_banding = False
    props = shape._element.graphic.graphicData.tbl.tblPr
    style = props.find(qn("a:tableStyleId"))
    if style is None:
        style = OxmlElement("a:tableStyleId")
        props.append(style)
    style.text = GRID_STYLE
    for index, height in enumerate(block.heights):
        table.rows[index].height = Emu(int(Inches(height)))
    return table


def _write(
    cell: Any,
    text: str,
    size: float,
    *,
    bold: bool = False,
    fill: RGBColor | None = None,
    color: RGBColor | None = None,
) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill or WHITE
    cell.margin_left = cell.margin_right = Inches(0.05)
    cell.margin_top = cell.margin_bottom = Inches(0.03)
    frame = cell.text_frame
    frame.word_wrap = True
    _paragraphs(frame, text, size, bold=bold, color=color)


def _box(
    page: Any,
    left: float,
    top: float,
    width: float,
    height: float,
    text: str,
    size: float,
    *,
    bold: bool = False,
    color: RGBColor | None = None,
) -> None:
    shape = page.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.word_wrap = True
    _paragraphs(frame, text, size, bold=bold, color=color)


def _paragraphs(
    frame: Any, text: str, size: float, *, bold: bool, color: RGBColor | None
) -> None:
    """줄마다 문단 하나. 한글 글꼴(`a:ea`)과 언어(`ko-KR`)를 박는다 — 안 박으면 PowerPoint 가
    한글을 영문 글꼴로 그리고, 줄바꿈을 단어가 아니라 글자에서 끊는다."""
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = line
        font = run.font
        font.size = Pt(round(size, 1))
        font.bold = bold
        font.name = FONT
        if color is not None:
            font.color.rgb = color
        props = run._r.get_or_add_rPr()
        props.set("lang", "ko-KR")
        props.set("altLang", "en-US")
        east = props.find(qn("a:ea"))
        if east is None:
            # `a:latin`(font.name) 뒤에 — rPr 의 차례(latin · ea)를 지켜야 PowerPoint 가 연다.
            east = OxmlElement("a:ea")
            props.append(east)
        east.set("typeface", FONT)
        # 빈 칸의 높이는 끝 문단 속성(`a:endParaRPr`)의 크기를 따른다 — 안 박으면 18pt 로 서서
        # 빈 칸이 있는 줄만 높아진다(첫 판에서 그랬다).
        end = paragraph._p.get_or_add_endParaRPr()
        end.set("sz", str(round(size * 100)))
        end.set("lang", "ko-KR")
